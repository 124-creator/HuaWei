# /// script
# requires-python = ">=3.12"
# dependencies = ["python-docx"]
# ///
"""Check whether the teammate draft shares our frozen R12 numbers and vocabulary."""
import json
import re
from pathlib import Path

from docx import Document

SRC = Path(r"C:\Users\Dreamboat\OneDrive\文档\WeChat Files\wxid_tnssm99dpsea22\FileStorage\Temp\Copy\2026华为杯A题_本人部分P0候选_v0.6.docx")
NUMBERS = ["3.7433", "3.7432", "2.0531", "2.053", "4.1311", "4.1310", "1.016766", "1.0168",
           "43.4760", "43.47", "34.4910", "34.49", "31.4081", "31.40", "1.027400", "1.0274",
           "4.1875", "4.2190", "4.2189", "23.4614", "10.6%", "10.6 %"]
WORDS = ["R12", "快速插入", "微批", "Treap", "依赖带", "安全链", "spill", "Spill", "SPILL",
         "同方案", "C3", "启发式投影", "两阶段", "官方评估器"]


def main() -> None:
    doc = Document(SRC)
    text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            text += "\n" + " | ".join(cell.text for cell in row.cells)
    hits = {n: text.count(n) for n in NUMBERS if n in text}
    words = {w: text.count(w) for w in WORDS if w in text}
    contexts = []
    for token in ("3.743", "4.131", "2.053"):
        for m in re.finditer(re.escape(token), text):
            contexts.append(text[max(0, m.start() - 70): m.end() + 40].replace("\n", " "))
    print(json.dumps({"number_hits": hits, "word_hits": words, "contexts": contexts[:8],
                      "total_chars": len(text)}, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
