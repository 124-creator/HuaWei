# /// script
# requires-python = ">=3.12"
# dependencies = ["python-docx"]
# ///
"""Pull the teammate draft's headline numbers (abstract + key table rows) for comparison."""
import json
from pathlib import Path

from docx import Document

SRC = Path(r"C:\Users\Dreamboat\OneDrive\文档\WeChat Files\wxid_tnssm99dpsea22\FileStorage\Temp\Copy\2026华为杯A题_本人部分P0候选_v0.6.docx")


def main() -> None:
    doc = Document(SRC)
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    abstract = next((t for t in paras if t.startswith("多核神经网络处理器")), None)
    result = {"abstract_head": abstract[:600] if abstract else None, "tables": []}
    for i, table in enumerate(doc.tables[:6]):
        rows = [" | ".join(cell.text.strip()[:22] for cell in row.cells) for row in table.rows[:4]]
        result["tables"].append({"index": i, "head": rows})
    print(json.dumps(result, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
