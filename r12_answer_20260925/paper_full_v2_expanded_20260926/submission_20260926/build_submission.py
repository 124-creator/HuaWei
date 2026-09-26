#!/usr/bin/env python3
# 本程序及代码是在人工智能工具辅助下完成的。
# ChatGPT，GPT-6 Astra Pro，OpenAI；Astra系列公告日期2026-09-03。
# 用途：冻结文稿编辑、只读核对与排版，不调用求解器或评估器。
from __future__ import annotations
import argparse, ast, copy, csv, hashlib, importlib.util, json, re, tempfile, zipfile
from pathlib import Path
from collections import Counter
from statistics import median
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph
from docx.table import Table
import pypandoc

HERE=Path(__file__).resolve().parent
V2=HERE.parent
BASE=V2.parent
SOURCE=V2/'out/R12论文_完整增强版_匿名.docx'
OUT=HERE/'out'
NAME='多核NPU切图调度论文'
EXPECTED='003f6d96515d45cdb024c7ffe8d6449c4d7744ba7e95f92fd7aabadc369a6322'
spec=importlib.util.spec_from_file_location('document_helpers',V2/'build_paper.py')
H=importlib.util.module_from_spec(spec);spec.loader.exec_module(H)
TOOL='ChatGPT（GPT-6 Astra Pro，OpenAI；Astra系列公告日期2026年9月3日）'

def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def alltext(e):return ''.join(e.xpath('.//w:t/text()|.//m:t/text()'))
def text_nodes(e):return e.xpath('.//w:t|.//m:t')
def patch_text(e,pattern,replacement):
    nodes=text_nodes(e);parts=[x.text or '' for x in nodes];s=''.join(parts)
    for m in reversed(list(re.finditer(pattern,s))):
        start,end=m.span();rep=replacement(m) if callable(replacement) else replacement
        a=0;hits=[]
        for node,part in zip(nodes,parts):
            b=a+len(part)
            if b>start and a<end:hits.append((node,max(0,start-a),min(len(part),end-a),part))
            a=b
        if not hits:continue
        if len(hits)==1:
            n,l,r,part=hits[0];n.text=(n.text or '')[:l]+rep+(n.text or '')[r:]
        else:
            n,l,r,part=hits[0];n.text=(n.text or '')[:l]+rep
            for node,_,_,_ in hits[1:-1]:node.text=''
            n,l,r,part=hits[-1];n.text=(n.text or '')[r:]
    return len(list(re.finditer(pattern,s)))

def rename(e):
    patch_text(e,r'R\s*12','LBSS');patch_text(e,r'C\s*3','SCB')
    patch_text(e,r'R6/R7/R10','基础构造、结构候选和反馈搜索')
    patch_text(e,r'A087','case_087（场景A）')
    patch_text(e,r'\bL2([1-5])\b',lambda m:'L2-'+m[1])
    patch_text(e,r'F28','图7-9');patch_text(e,r'F27','图7-8')
    patch_text(e,r'开发诊断案例','设计阶段的诊断案例')
    for old,new in [('本轮','本文'),('本稿','本文'),('原存档','归档'),('原有','已有'),('读表统计','记录汇总'),('待核验','核对')]:
        patch_text(e,re.escape(old),new)
    patch_text(e,r'(?<!\d)(4\.8|4\.9|7\.6|7\.7|8\.7|8\.8)(?!\d)',lambda m:{'4.8':'4.9','4.9':'4.8','7.6':'7.7','7.7':'7.6','8.7':'8.8','8.8':'8.7'}[m[0]])

def fill(p,text,kind=None):
    p.clear();p.add_run(text)
    if kind is None:
        sn=p.style.name
        kind='h'+sn[-1] if sn.startswith('Heading ') else ('caption' if sn in ('图题','表题') else 'body')
    H.para_format(p,kind)
    return p

def add_after(d,anchor,text,kind='body'):
    el=OxmlElement('w:p');anchor.addnext(el);p=Paragraph(el,d._body)
    if kind.startswith('h'):p.style=d.styles['Heading '+kind[-1]]
    elif kind=='caption':p.style=d.styles['表题']
    fill(p,text,kind);return p

def table_after(d,anchor,title,headers,rows,widths):
    cp=add_after(d,anchor,title,'caption');cp.paragraph_format.keep_with_next=True
    t=d.add_table(rows=1,cols=len(headers));cp._p.addnext(t._tbl)
    for c,s in zip(t.rows[0].cells,headers):c.text=str(s)
    for row in rows:
        for c,s in zip(t.add_row().cells,row):c.text=str(s)
    H.style_table(t,widths)
    return t

def stats():
    vf=BASE/'publication/sources/closeout/results/verification.json'
    q1=BASE/'publication/tables/Q1_per_case.csv'
    checks=json.loads(vf.read_text())['checks'];selected=[x for x in checks if x['kind']=='selected']
    assert len(selected)==1400
    rows=[]
    for sc,n,expected in [('A',400,70),('B',500,95),('L2',500,98)]:
        z=[x for x in selected if x['scene']==sc];assert len(z)==n
        c=sum(x['compute_gap_upper_bound']<=.03 for x in z);assert c==expected
        gaps=[x['compute_gap_upper_bound'] for x in z]
        ordered=sorted(gaps);q=lambda f:ordered[max(0,__import__('math').ceil(len(ordered)*f)-1)]
        med5=median(x['compute_gap_upper_bound'] for x in z if x['cores']==5)
        rows.append(dict(scene=sc,total=n,certificates=c,coverage=c/n,gap_min=min(gaps),gap_q25=q(.25),gap_median=median(gaps),gap_q75=q(.75),gap_max=max(gaps),five_core_median=med5))
    sums={}
    with q1.open(encoding='utf-8-sig') as f:
        for r in csv.DictReader(f):sums[r['case']]=sums.get(r['case'],0)+float(r['wall_seconds'])
    assert len(sums)==100 and sum(v>600 for v in sums.values())==17 and round(median(sums.values()),1)==90.5
    return {'derivation':'冻结表只读复算；未调用求解器或官方评估器','tool':TOOL,'certificate_field':'compute_gap_upper_bound based on global_compute_lower_bound; post-hoc, not runtime gate frequency','selected_configurations':1400,'certificate_count':263,'coverage':263/1400,'by_scene':rows,'A_wall_seconds':{'scope':'Sum successful archived requests for cores 2-5 per graph; includes successful extended-budget runs, excludes unrecorded failed-attempt elapsed time; not concurrent batch duration','median':median(sums.values()),'above_600_count':17,'per_case':sums},'evaluation_static_upper_bounds':{'A':23,'B':25,'L2':42,'scope':'Complete multicore official score calls per request, excluding Spill diagnostics, single-core REF and post-hoc pair scoring','A_terms':[4,6,4,2,6,1],'B_terms':[6,6,4,6,1,2],'L2_terms':[6,6,4,6,16,1,1,2]},'sources':{'verification.json':sha(vf),'Q1_per_case.csv':sha(q1)}}

EQ={
'4-6':r'\mathop{\mathrm{lex\ min}}_{P\in\mathcal F}\bigl(T(P),D_{\mathrm{added}}(P)\bigr)',
'4-10':r'f_u\le s_v\ \lor\ f_v\le s_u,\quad u\ne v,\ \mathrm{core}(u)=\mathrm{core}(v),\ \mathrm{pipe}(u)=\mathrm{pipe}(v)',
'4-11':r'r_s^{\mathrm{seq}}=\begin{cases}f_{\mathrm{prev}(s)}+100,&\mathrm{prev}(s)\ne\varnothing,\\0,&\mathrm{prev}(s)=\varnothing.\end{cases}',
'4-12':r'r_s^{\mathrm{dep}}=\max\bigl(\{0\}\cup\{f_p+1000\,\mathbf1_{a(p)\ne a(s)}:p\in\mathrm{Pred}(s)\}\bigr)',
'4-13':r'r_s=\max\{r_s^{\mathrm{seq}},r_s^{\mathrm{dep}}\}',
'5-5':r'\widehat r_c(s)=\max\left\{\ell_c+100\,\mathbf1_{\pi_c\ne\varnothing},\ \max\bigl(\{0\}\cup\{\widehat f_p+1000\,\mathbf1_{a(p)\ne c}:p\in\mathrm{Pred}(s)\}\bigr)\right\}',
'5-6':r'\mathrm{rank}(u)=d_u+\max\left(\{0\}\cup\left\{\delta+\frac{2b_{uv}}{60}+\mathrm{rank}(v):v\in\mathrm{Succ}(u)\right\}\right)',
'5-8':r'\begin{aligned}\mathrm{span}+\varepsilon&<\mathrm{need}\quad\lor\\\mathrm{rightmost}+\varepsilon&<\mathrm{ready}+\mathrm{need}.\end{aligned}',
'6-2':r'D^{B}_{\mathrm{boundary}}(t)=\begin{cases}2rb,&r>0,\\0,&r=0.\end{cases}',
'6-5':r'w\in\{4,8\},\qquad |\mathcal W|\le16,\qquad |U_W|\le128\quad(W\in\mathcal W)',
'6-10':r'N_j=\left|\{i:T_{\mathrm{stage}\,j+1,i}<T_{\mathrm{stage}\,j,i}\}\right|,\qquad(N_{\mathrm{ins}},N_{\mathrm{mb}})=(12,2)',
'8-2':r'\eta=\frac{\max_c W_c}{\left(\sum_c W_c\right)/k}',
'8-3':r'p=\min\left(1,\ 2^{1-n}\sum_{j=0}^{m}\frac{n!}{j!(n-j)!}\right)'
}

def math_blocks():
    md='\n\n'.join('FORMULA'+str(i)+'\n\n$$'+f+'$$' for i,f in enumerate(EQ.values()))
    with tempfile.TemporaryDirectory() as td:
        p=Path(td)/'math.docx';pypandoc.convert_text(md,'docx',format='markdown',outputfile=str(p))
        d=Document(p);out={};current=None
        for p in d.paragraphs:
            if p.text.startswith('FORMULA'):current=int(p.text[7:])
            ms=p._p.xpath('.//m:oMath')
            if ms and current is not None:out[list(EQ)[current]]=copy.deepcopy(ms[0]);current=None
    assert len(out)==len(EQ)
    for el in out.values():
        for t in list(el.xpath('.//m:t')):
            if t.text=='&':r=t.getparent();r.getparent().remove(r)
    return out

def move_range(d,start,end,before):
    es=list(d.element.body);a=es.index(start);b=es.index(end)
    for el in es[a:b]:before.addprevious(el)

def algorithm(d,ps,head,inp,body,last,lines):
    before=ps[head]._p;t=d.add_table(rows=1,cols=1);before.addprevious(t._tbl)
    cell=t.cell(0,0);cell.text='';content=[ps[head].text,ps[inp].text,*lines]
    for i,text in enumerate(content):
        p=cell.paragraphs[0] if i==0 else cell.add_paragraph();fill(p,text,'algo')
        p.paragraph_format.keep_with_next=(i<len(content)-1)
        for r in p.runs:H.set_font(r,12,i==0)
        if i>1:p.paragraph_format.left_indent=Pt(8)
    H.widths(t,[16.5]);H.borders(t,True)
    t.rows[0]._tr.get_or_add_trPr().append(H.node('w:cantSplit'))
    for i in (head,inp,body):ps[i]._p.getparent().remove(ps[i]._p)
    fill(ps[last],'各阶段共用预算Γ；近优门控仅跳过其覆盖的控制阶段，外层插入和微批按剩余预算执行。')


def prepare():
    assert sha(SOURCE)==EXPECTED,'Source drift: stop before rebuilding.'
    d=Document(SOURCE);ps=list(d.paragraphs);ts=list(d.tables)
    assert len(ps)==792 and len(d.inline_shapes)==37 and len(ts)==110
    archive=[[[c.text for c in r.cells] for r in t.rows] for t in ts[100:109]]
    numeric_before={i:re.findall(r'(?<![A-Za-z])[-+]?\d+(?:\.\d+)?',alltext(t._tbl)) for i,t in enumerate(ts) if not (len(t.rows)==1 and len(t.columns)==2)}
    edits={int(k):v for k,v in json.loads((HERE/'edits.json').read_text())['paragraphs'].items()}
    # Reconcile prose anchors against the pinned document before applying edits.
    edits[256]=edits.pop(258);edits[313]=edits.pop(316);edits[627]=edits.pop(628);edits.pop(231)
    extra={
      229:'本章以图结构构造有限候选，按第4章的下界筛选与官方评价规则选择方案，保留400个配置的完整性能、搬运和预算记录。',
      240:'两端独占是本文采用的链收缩充分条件。若收缩后形成经过新块的有向环，则原图中存在绕过该边连接其两端的路径，与源端唯一后继或目标端唯一前驱矛盾，因此商图保持无环。性能变化仍由完整评价确定。',
      261:'设浮点容差ε=16 ulp(·)，其中ulp按原实现的相应边界数值计算，need=duration+gap。式(5-8)筛除不可能容纳任务的子树，首个可行位置仍按式(5-7)判定；图5-3展示候选日历查询。',
      264:'场景A不启用当前驻留型微批，其候选侧重独立Task的粒度和释放次序；B与L2中的微批用于调整同核驻留。',
      272:'近优门控作用于控制阶段；外层插入继续按剩余预算执行，失败状态与成功方案分开记录。',
      273:'Task激活下界将固定方案的乐观Task时长代入释放约束，用于筛选当前候选。实现中的全局下界断言与归档校验结果一并列于支撑材料。',
      294:'题面要求的1至5核平均加速比见图5-5，一核点取1。两核至五核均值由1.855升至3.743，相邻增量分别为0.705、0.637、0.546；每点均由100张图的比值取算术平均。',
      298:'LBSS与SCB的平均性能差距随核数扩大，图5-5给出两条完整曲线。',
      303:'注：同一输入和官方评价口径下的逐图比较，时间单位为cycles；方法与预算说明见第4.3.3节。',
      304:'五核逐图耗时降幅均值为43.48%，对应LBSS完整流程的最终结果。',
      314:'图5-6给出全部400个配置的对数比值，并逐项定位9个退化配置。',
      326:'注：两例均比SCB更快且新增COPY更高，表中列出最终方案值。',
      434:'注：阶段结果来自同一总预算下的顺序执行；未改善包括持平、未触发和未完成。',
      486:'单次搬运的理想带宽时间相差约4.17倍，整图结果还由计算、同步、流水线和其他搬运共同决定。',
      499:'右侧第二项为同一L2配置下更换方案的选解比，三项构成逐例恒等式。',
      550:'500对中8对退化，最差为case_026五核的0.993980，约慢0.6%；两核有4条，case_092在三核与五核各出现1条。图7-6放大显示全部8条记录。',
      579:'注：SCB与LBSS均在L2配置下评价；比较定义见第4.3.3节。',
      581:'五核L2相对SCB的逐图降幅为31.41%，场景B为34.49%。各场景采用对应的官方评价与实际选解，固定方案的配置效应另由第7.2节给出。',
      586:'本节汇总测试覆盖与框架复用范围。',
      588:'100张正式图覆盖全部指定核数与场景，共1400个配置；统计按场景、核数分组，并在同图间配对。',
      591:'五核三场景最小加速比为1.202、1.812、1.932。全部1400个最终结果不慢于REF，其中场景A包含case_087的1800秒补跑。',
      592:'完整分布与预算例外共同描述不同实例的收益范围。',
      594:'三问复用方案表示、合法性检查、下界与择优。场景A配置依赖带与反馈分核，B/L2配置关键链移动和微批，L2另配置跨场景重评价；表7-9列出启用范围。',
      598:'当前驻留型微批只在B与L2中启用。',
      600:'注：●表示启用，○表示未启用；按实际求解流程整理。',
      601:'迁移到其他平台时，应重新核对通信、容量、同步、缓存及评价规则，并检验资源下界的适用条件。',
      603:'外推范围由图结构、硬件配置和请求预算共同限定。',
      608:'代理量用于快速构造，下界用于安全筛选，官方值用于最终择优。这三个环节保持独立职责。',
      609:'可复用的是构造、筛选与评价的分工；具体规则按平台资源确定。',
      612:'注：常规单请求软预算600秒，case_087场景A四个配置另列1800秒补跑；时间列为实际请求墙钟。',
      618:'本章评价固定为1 MiB、250 bytes/cycle，完整对照与逐例记录见附录A。',
      628:'算术均值使用以n为分母的经验协方差，有如下关系：',
      671:'表8-5按最终方案首次所属组件记录来源。',
      686:'表8-6 相邻核数扩展的逐例结果（每行100图）'
    }
    edits.update(extra)
    for i,text in edits.items():
        old=alltext(ps[i]._p)
        assert not ps[i]._p.xpath('.//w:drawing'),f'Cannot overwrite a figure at paragraph {i}'
        if old.startswith('图') and re.match(r'图\d',old):assert text.startswith('图'),f'Caption anchor mismatch {i}'
        fill(ps[i],text)
    S=stats()
    # Move whole sections, including their formula/table nodes, before each chapter summary.
    move_range(d,ps[210]._p,ps[224]._p,ps[206]._p)
    move_range(d,ps[619]._p,ps[632]._p,ps[610]._p)
    move_range(d,ps[684]._p,ps[690]._p,ps[682]._p)
    coverrows=[[z['scene'],z['total'],z['certificates'],f"{100*z['coverage']:.1f}",f"{100*z['five_core_median']:.1f}"] for z in S['by_scene']]
    coverrows.append(['合计',1400,263,'18.8','—'])
    table_after(d,ps[223]._p,'表4-4 全局计算下界的事后证书覆盖',['场景','配置数','3%证书数','覆盖率/%','五核中位间隙上界/%'],coverrows,[2,2.4,3,3,6.1])
    table_after(d,ps[88]._p,'表3-3 缩写与补充符号',['记号','含义'],[
      ['LBSS','下界引导的结构化搜索'],['SCB','静态构造基线'],['REF','题目规定的整图单核参考'],['Task','一次下发的执行任务'],['Spill','容量不足触发的缓存换出与换入'],['Treap','用于候选日历查询的树堆区间索引'],['FIFO','先进先出缓存策略'],['Γ','一次请求的总软预算，单位秒'],['δ','插入候选中的同步延迟代理，单位周期'],['η','各核静态计算负载不均衡度'],['ℓᵢ、Δᵢ','配对对数比与单例时间降幅，分别定义']], [3.1,13.4])
    table_after(d,ps[468]._p,'表7-13 题面要求与本文对应内容',['题面要求','模型或机制','正文位置'],[
      ['容量影响','单张量可缓存性、FIFO淘汰、在途读取与复用次序','7.1.2、7.1.4'],['带宽影响','独立DDR/L2服务池及乐观资源下界','7.1.3、式(4-25)'],['数学模型','FIFO容量、字节命中率、双池约束','式(7-1)至式(7-4)'],['求解算法','B结构候选、跨配置重评价、插入与微批','7.2.4、算法7-1'],['结果与附件','1至5核曲线、配置加速比及逐例指标','7.3、附录A.3']], [3,8.7,4.8])
    # One definition follows the sign-test formula; Wilcoxon starts in its own paragraph.
    add_after(d,ps[652]._p,'Wilcoxon符号秩检验[19]还使用差异大小：定义ℓᵢ=ln(T_SCB,i/T_LBSS,i)，剔除零差异后按|ℓᵢ|计算平均秩rkᵢ。')
    patch_text(ps[259]._p,'σ','δ')
    patch_text(ps[631]._p,r'd(?=i|ᵢ|=)|d(?=̄)','Δ')
    for i in (390,508):patch_text(ps[i]._p,r'预算\s*B','预算Γ')
    # Appendix D: formal, source-supported disclosure; unknown historical identifiers remain explicit.
    for i in range(786,792):
        if ps[i]._p.getparent() is not None:ps[i]._p.getparent().remove(ps[i]._p)
    fill(ps[785],'附录D 人工智能工具使用声明','h2')
    d1=add_after(d,ps[785]._p,'本文在资料检索、文稿整理、辅助编程、统计核对和排版中使用人工智能工具。工具及采用范围如下；模型论证、归档结果与工具输出分别核对，人工智能工具不作为实验数据来源。')
    td=table_after(d,d1._p,'表D-1 工具、版本信息与使用范围',['工具与记录型号','开发机构及发布日期','使用环节'],[
      ['Claude / Claude Code；历史提交记录标注Claude Opus 5.5','Anthropic；2026-09-22（模型发布日期）','初稿相关章节、历次文稿修改、统计与排版辅助；模拟审稿意见'],
      ['ChatGPT；本次GPT-6 Astra Pro','OpenAI；2026-09-03（Astra系列公告，Pro未另列发布日期）','审稿意见处理、冻结表只读复算、公式与术语核对、图中文字整理、程序注释与排版']], [5.2,5.2,6.1])
    p=add_after(d,td._tbl,'《模拟审稿意见》由Claude Code（Anthropic）模拟生成，属于AI辅助审阅，不是组委会或真人评审意见。其内容在本次文稿修改中被参考，因此一并披露。历史提交中的型号标签用于说明可追溯记录，不能据此补定所有历史会话的后台版本。')
    p=add_after(d,p._p,'历史ChatGPT会话的完整后台型号、Claude Code客户端版本及部分早期开发环节未在现有材料中完整留存。本声明如实列出能够确认的记录，不将本次型号替代全部历史型号。模型发布日期依据开发机构公告，客户端版本与模型版本分开理解。')
    p=add_after(d,p._p,'本次优化仅对已保存的文稿、图件和结果进行编辑与汇总，未运行原始计算图的求解器或官方评估器。新增的下界覆盖、间隙分布和按图合计墙钟来自冻结记录；数据来源、分母、单位及脚本校验值列于支撑材料。')
    p=add_after(d,p._p,'数学模型按题面规则建立，命题附推导，引用的算法与统计方法在正文标注来源。自编提交程序添加AI辅助注释，官方评估程序保持原样。AI生成的建议不等同于正确结论，最终文稿、程序与数据的理解、采用和责任仍由参赛队承担。')
    p=add_after(d,p._p,'模型与日期来源：Anthropic的Claude Opus 5.5公告（2026-09-22）；OpenAI的GPT-6 Astra及安全概览公告（2026-09-03）。相关可查询地址列入随附工具信息记录。')
    # Replace the internal paths with the actual support-package structure.
    support=[['内容','包内目录','说明'],['统一程序入口','program/solve.py','仅整理调用入口，不改变算法'],['自编实现与原版评价器','program/engine/','前者补AI注释；后者逐字节保留'],['固定配置','data/config.txt','与官方配置一致'],['逐例与汇总结果','results/tables/','保留全部原表'],['校验记录','results/verification.json','1400配置及同方案对照核查'],['只读统计','analysis/recompute.py','不导入求解器或评估器'],['文件清单','manifest.json','列出实际提供文件和仍需的运行输入']]
    for r,values in zip(ts[109].rows,support):
        for c,v in zip(r.cells,values):c.text=v
    H.style_table(ts[109],[4,6,6.5])
    # These are not observations of a worse case: zero-regression rows display not-applicable.
    table98_before=[[c.text for c in r.cells] for r in ts[98].rows]
    for ri,r in enumerate(ts[98].rows[1:],1):
        c=r.cells[-1]
        if c.text and c.text!='—' and not c.text.startswith('case_'):c.text='case_'+c.text.replace('case_','')
        if ri in (1,5,8):r.cells[-2].text='—';r.cells[-1].text='—'
    for pi in [51,223,335,621,633,681,685]:
        ps[pi].paragraph_format.space_before=Pt(3);ps[pi].paragraph_format.space_after=Pt(3)
    # Formal algorithm boxes; retain conditional control and the stated numerical limits.
    algorithm(d,ps,269,270,271,272,[
      '1 由原图构造基础候选，完整评价后保存首个成功方案。',
      '2 依次处理减少活动核心、依赖带、反馈重排与分量家族。',
      '3   若达到3%近优条件，跳过受门控阶段；否则生成该阶段候选。',
      '4   检查覆盖与联合无环，去重；固定方案下界严格大于当前时间时剪枝。',
      '5   在剩余预算内完整评价，按（Makespan，新增COPY，名称）择优。',
      '6 控制阶段结束后，预算允许时生成Treap插入候选并按第4至5步处理。',
      '7 输出两字段方案与报告；无成功方案则记录失败。'])
    algorithm(d,ps,389,390,391,392,[
      '1 生成基础候选、减少活动核心候选及关键链移动，保存成功解。',
      '2 未满足近优条件时，生成驻留组织候选，剪枝后至多评价6个。',
      '3 构造插入候选并完整评价，按字典序更新已得方案。',
      '4 若当前Spill新增>0且有剩余预算，诊断大张量换入换出事件。',
      '5   按式(6-4)、式(6-5)生成至多2个微批候选，检查并评价。',
      '6 输出成功方案及诊断记录。生成30秒、插入评价120秒、诊断40秒、',
      '  单个微批评价45秒，均计入600秒总软预算。'])
    algorithm(d,ps,507,508,509,510,[
      '1 在L2配置下生成基础、减少活动核心及关键链移动候选。',
      '2 未满足近优条件时，处理分量家族，至多评价6个新候选。',
      '3 若预算允许且仍未近优，内部求解B方案至多180秒，再作L2评价至多40秒。',
      '4 生成插入候选至多30秒，评价至多120秒。',
      '5 若当前Spill新增>0，诊断并处理至多2个局部微批候选。',
      '6 各步完整成功且字典序更优时更新方案；输出最终L2方案。'])
    # Shared public terminology, chapter cross-references and math subscripts.
    for p in d.element.body.xpath('.//w:p'):rename(p)
    patch_text(d.element.body,r'\bround12\b','算法实现')
    # Upright operators and corrected equivalent display forms.
    maths=math_blocks()
    for t in d.tables:
        if len(t.rows)==1 and len(t.columns)==2:
            key=t.cell(0,1).text.strip('()')
            if key in maths:
                p=t.cell(0,0).paragraphs[0];p.clear();p._p.append(copy.deepcopy(maths[key]));H.para_format(p,'eq')
    # Operator names already split across OMML runs are kept upright as a group.
    for m in d.element.body.xpath('.//m:oMath'):
        for run in m.xpath('.//m:r'):
            text=''.join(run.xpath('.//m:t/text()'))
            if text in {'max','min','Pred','Succ','rank','span','need','ready','rightmost','duration','gap','seq','dep','LBSS','SCB','lex','core','pipe'}:
                rp=run.find(qn('m:rPr'))
                if rp is None:rp=OxmlElement('m:rPr');run.insert(0,rp)
                rp.append(OxmlElement('m:nor'))
    # Update table numbers by first occurrence after the section moves and new table insertions.
    ids=[];counts={};mapping={}
    for p in d.paragraphs:
        m=re.match(r'^表\s*([3-8])-([0-9]+)\s',p.text)
        if m:
            old=f'{m[1]}-{m[2]}'
            if old not in mapping:
                counts[m[1]]=counts.get(m[1],0)+1;mapping[old]=f'{m[1]}-{counts[m[1]]}'
    for p in d.element.body.xpath('.//w:p'):
        patch_text(p,r'表\s*([3-8]-[0-9]+)(?![0-9])',lambda m:'表'+mapping.get(m[1],m[1]))
    # Abstract-page template wording; anonymous body follows the explicit format prohibition.
    for text in ['中国研究生创新实践系列大赛','“华为杯”第二十三届中国研究生','数学建模竞赛']:
        p=ps[0].insert_paragraph_before(text);H.para_format(p,'caption')
        p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(2)
        for r in p.runs:H.set_font(r,12,True,'宋体')
    fill(ps[0],'题 目：'+NAME.replace('论文','') if False else '题 目：基于可证下界与结构化搜索的多核NPU切图调度','h1')
    for r in ps[0].runs:H.set_font(r,16,True,'黑体')
    ps[0].paragraph_format.space_before=Pt(8);ps[0].paragraph_format.space_after=Pt(8)
    fill(ps[1],'摘 要：','body');ps[1].paragraph_format.first_line_indent=Pt(0)
    for r in ps[1].runs:r.bold=True
    ps[9].paragraph_format.page_break_before=True
    # Black commands, no stretched paths; make all text and metadata anonymous.
    for i in [772,773]:
        if ps[i]._p.getparent() is not None:
            H.para_format(ps[i],'code')
            for r in ps[i].runs:
                H.set_font(r,12);r.font.name='Consolas';r.font.color.rgb=RGBColor(0,0,0)
    for p in d.element.body.xpath('.//w:p'):
        for c in p.xpath('.//w:color'):c.set(qn('w:val'),'000000')
    for el in d.element.body.xpath('.//wp:docPr|.//pic:cNvPr'):
        for key in ['descr','title']:el.attrib.pop(key,None)
        if 'name' in el.attrib:el.set('name','论文图件')
    cp=d.core_properties
    cp.title='基于可证下界与结构化搜索的多核NPU切图调度'
    cp.subject='多核调度、结构化搜索与缓存分析';cp.author='';cp.last_modified_by='';cp.comments='';cp.keywords='LBSS; SCB; NPU'
    # Preserve existing data tables. Only the non-applicable display in Table 8-6 is changed.
    assert [[[c.text for c in r.cells] for r in t.rows] for t in ts[100:109]]==archive
    assert sum(len(t.rows)-1 for t in ts[100:109])==900
    equations=[t.cell(0,1).text for t in d.tables if len(t.rows)==1 and len(t.columns)==2 and re.fullmatch(r'\(\d+-\d+\)',t.cell(0,1).text)]
    assert len(equations)==64 and len(set(equations))==64
    text=alltext(d.element.body)
    bad={k:len(re.findall(k,text)) for k in [r'R12',r'C3',r'R6',r'R7',r'R10',r'round12',r'F27',r'F28',r'A087',r'\bL2[1-5]\b',r'/home/',r'paper_full_v1/',r'publication/sources/',r'figures_v3/']}
    assert not any(bad.values()),bad
    for word in ['下界引导的结构化搜索','静态构造基线','263','18.8%','90.5','17张图','模拟审稿意见']:
        assert word in text,word
    negative_pattern=r'不|不能|并非|未|无法|尚未'
    oldcount=len(re.findall(negative_pattern,alltext(Document(SOURCE).element.body)))
    newcount=len(re.findall(negative_pattern,text))
    return d,ps,ts,S,{'source_sha256':EXPECTED,'prose_edits':len(edits),'archive_rows':900,'equations':64,'figures':37,'public_internal_tokens':bad,'negative_expression_token_count':{'pattern':negative_pattern,'before':oldcount,'after':newcount,'note':'Token count, not an independent linguistic rating; formal limitations retained'},'table_number_map':mapping,'table8_6_before':table98_before,'frozen_data_changed':False,'new_statistics_scope':S['derivation'],'official_program_execution':False,'formal_word_font_export':'NOT_COMPLETED; Linux preview is not a Microsoft Word font-conformant export','AI_history_metadata':'PARTIAL; historical ChatGPT model and some client records unavailable','source_review':'模拟审稿意见.md, blob 38618b12e85c56882f5238ad9361972854d04c63; AI-generated by Claude Code'},archive

def preflight():
    d,ps,ts,S,receipt,archive=prepare();dest=HERE/'preflight';dest.mkdir(exist_ok=True)
    (dest/'preflight.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2))
    (dest/'statistics.json').write_text(json.dumps(S,ensure_ascii=False,indent=2))
    # Small fragment documents only. The complete paper is not saved here.
    a=Document();a.sections[0].left_margin=Cm(2.25);a.sections[0].right_margin=Cm(2.25);a.sections[0].top_margin=Cm(2.5);a.sections[0].bottom_margin=Cm(2)
    for el in d.element.body:
        if el is ps[9]._p:break
        a.element.body.insert(-1,copy.deepcopy(el))
    a.save(dest/'abstract.docx')
    m=Document();m.sections[0].left_margin=Cm(2.25);m.sections[0].right_margin=Cm(2.25)
    for t in d.tables:
        if len(t.rows)==1 and len(t.columns)==2 and t.cell(0,1).text.strip('()') in EQ:m.element.body.insert(-1,copy.deepcopy(t._tbl))
    m.save(dest/'equations.docx')
    print(json.dumps(receipt,ensure_ascii=False,indent=2))

def build():
    if (OUT/(NAME+'.docx')).exists():raise RuntimeError('Single-build guard: target manuscript already exists.')
    d,ps,ts,S,receipt,archive=prepare()
    from figures import replace_figures
    figs=replace_figures(d,BASE,HERE/'assets');receipt['figure_relabelling']=figs
    OUT.mkdir(exist_ok=True)
    d.save(OUT/(NAME+'.docx'))
    # Verify the serialized artifact, not only the in-memory table objects.
    saved=Document(OUT/(NAME+'.docx'))
    tables=[t for t in saved.tables if len(t.rows)==101]
    assert len(tables)==9
    assert [[[c.text for c in r.cells] for r in t.rows] for t in tables]==archive
    receipt.update(word_full_build_count=1,word_sha256=sha(OUT/(NAME+'.docx')),serialized_archive_rows=900,status='CONTENT_VALIDATED; preview rendering and formal Word export tracked separately')
    (OUT/'新校验记录.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2))
    (OUT/'补报统计.json').write_text(json.dumps(S,ensure_ascii=False,indent=2))
    from support import package_support
    receipt['support']=package_support(BASE,HERE,S)
    (OUT/'新校验记录.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2))
    print('Full manuscript built exactly once:',OUT/(NAME+'.docx'))

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--mode',choices=['preflight','build'],required=True);a=ap.parse_args()
    preflight() if a.mode=='preflight' else build()
