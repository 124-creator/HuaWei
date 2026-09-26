#!/usr/bin/env python3
"""Fix cross-renderer OMML operators; preserve equations and all frozen values.
本程序及代码是在人工智能工具辅助下完成的。ChatGPT / OpenAI；
具体版本与发布日期须由参赛队据实核验。
"""
from pathlib import Path
from copy import deepcopy
import tempfile
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pypandoc
from build_paper import para_format, set_font
ROOT=Path(__file__).resolve().parent
NAME='R12论文_完整增强版_匿名'
def apply(path):
    d=Document(path)
    removed=stars=0
    for t in list(d.element.body.xpath('.//m:t')):
        if t.text=='&':
            r=t.getparent();r.getparent().remove(r);removed+=1
        elif t.text=='*' or (t.text in ('+','−','-') and t.getparent().getparent().tag in (qn('m:sub'),qn('m:sup')) and len(t.getparent().getparent())==1):
            if t.text=='*':t.text='∗'
            r=t.getparent();pr=r.find(qn('m:rPr'))
            if pr is None:pr=OxmlElement('m:rPr');r.insert(0,pr)
            pr.append(OxmlElement('m:nor'));stars+=1
    formula=r'p=\mathrm{min}\left(1,\ 2^{1-n}\sum_{j=0}^{m}\frac{n!}{j!(n-j)!}\right)'
    with tempfile.TemporaryDirectory() as td:
        tmp=Path(td)/'eq.docx'
        pypandoc.convert_text('$$'+formula+'$$','docx',format='markdown',outputfile=str(tmp))
        math=deepcopy(Document(tmp).element.body.xpath('.//m:oMath')[0])
    for table in d.tables:
        if len(table.rows)==1 and len(table.columns)==2 and table.cell(0,1).text=='(8-3)':
            p=table.cell(0,0).paragraphs[0]
            for e in list(p._p):
                if e.tag!=qn('w:pPr'):p._p.remove(e)
            p._p.append(math);para_format(p,'eq')
            parent=table._tbl.getnext();target=None
            for p2 in d.paragraphs:
                if p2._p is parent:target=p2;break
            if target is not None:
                r=target.add_run(' 式(8-3)中 n=n₊+n₋，m 为 n₊ 与 n₋ 的较小者。');set_font(r,12)
            else:
                from docx.text.paragraph import Paragraph
                e=OxmlElement('w:p');table._tbl.addnext(e);pp=Paragraph(e,d._body)
                pp.add_run('式(8-3)中 n 为非零差异数，m 为更快与更慢图数的较小者，即 n=n₊+n₋、m=min(n₊,n₋)。');para_format(pp,'body')
            break
    for color in d.element.body.xpath('.//w:color'):color.set(qn('w:val'),'000000')
    d.save(path)
    print('OMML normalization:',removed,'alignment markers;',stars,'literal operators; factorial sign-test form.')
if __name__=='__main__':apply(ROOT/'out'/(NAME+'.docx'))
