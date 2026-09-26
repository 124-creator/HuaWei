#!/usr/bin/env python3
"""Read-only R12 paper expansion; original solver, graphs and scores are never run or edited.
本程序及代码是在人工智能工具辅助下完成的。工具：ChatGPT，OpenAI；
具体模型版本与发布日期由参赛队依据真实会话记录核验补全。
"""
from __future__ import annotations
import argparse, copy, hashlib, json, math, re, tempfile, zipfile, io
from pathlib import Path
from collections import Counter
from statistics import mean, median
import numpy as np
from scipy.stats import spearmanr
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
import pypandoc
from PIL import Image

HERE=Path(__file__).resolve().parent
EXPECTED='64ed65d80b63e7e7e2bbc5bbe13f22988f1c4ed25ecb2ac2e6018e7b695a724d'
NAME='R12论文_完整增强版_匿名'
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
WIDTH=16.5

def txt(el): return ''.join(el.xpath('.//w:t/text()'))
def node(tag,**attrs):
 x=OxmlElement(tag)
 for k,v in attrs.items():x.set(qn(k),str(v))
 return x

def set_font(run,size=12,bold=None,east='宋体'):
 run.font.name='Times New Roman';run.font.size=Pt(size)
 if bold is not None:run.bold=bold
 pr=run._r.get_or_add_rPr();fs=pr.find(qn('w:rFonts'))
 if fs is None:fs=node('w:rFonts');pr.insert(0,fs)
 for attr in ('ascii','hAnsi','cs'):fs.set(qn('w:'+attr),'Times New Roman')
 fs.set(qn('w:eastAsia'),east)
 for attr in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
  if qn('w:'+attr) in fs.attrib:del fs.attrib[qn('w:'+attr)]

def para_format(p,kind='body'):
 pf=p.paragraph_format
 pf.line_spacing=1.0;pf.space_before=Pt(0);pf.space_after=Pt(0)
 pf.left_indent=Pt(0);pf.right_indent=Pt(0);pf.first_line_indent=Pt(24)
 pf.keep_together=False;pf.keep_with_next=False;pf.widow_control=True;pf.page_break_before=False
 p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 pr=p._p.get_or_add_pPr()
 for tag in ['w:tabs','w:pBdr','w:shd','w:sectPr','w:numPr']:
  for x in list(pr.findall(qn(tag))):pr.remove(x)
 for x in list(pr.findall(qn('w:snapToGrid'))):pr.remove(x)
 pr.append(node('w:snapToGrid',**{'w:val':'0'}))
 if kind=='h1':
  pf.space_before=Pt(12);pf.space_after=Pt(8);pf.first_line_indent=Pt(0);pf.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 elif kind in ('h2','h3','h4'):
  pf.space_before=Pt(6);pf.space_after=Pt(3);pf.first_line_indent=Pt(0);pf.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.LEFT
 elif kind in ('caption','table','eq','algo','note','code'):
  pf.first_line_indent=Pt(0)
  if kind=='caption':p.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.space_before=Pt(3);pf.space_after=Pt(5);pf.keep_together=True
  elif kind=='table':p.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.space_after=Pt(0);pf.space_before=Pt(0);pf.keep_together=True
  elif kind=='eq':p.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.space_before=Pt(4);pf.space_after=Pt(4);pf.keep_together=True
  elif kind=='algo':p.alignment=WD_ALIGN_PARAGRAPH.LEFT;pf.space_before=Pt(2);pf.space_after=Pt(2);pf.keep_together=False
  elif kind=='code':p.alignment=WD_ALIGN_PARAGRAPH.LEFT;pf.space_after=Pt(4)
 for r in p.runs:set_font(r,14 if kind=='h1' else 12,kind in ('h1','h2','h3','h4','caption'),'黑体' if kind=='h1' else '宋体')
 for mr in p._p.xpath('.//m:r'):
  rp=mr.find(qn('w:rPr'))
  if rp is None:rp=node('w:rPr');mr.append(rp)
  rf=rp.find(qn('w:rFonts'))
  if rf is None:rf=node('w:rFonts');rp.insert(0,rf)
  rf.set(qn('w:ascii'),'Cambria Math');rf.set(qn('w:hAnsi'),'Cambria Math');rf.set(qn('w:eastAsia'),'宋体')
  for sz in ('w:sz','w:szCs'):
   z=rp.find(qn(sz))
   if z is None:z=node(sz);rp.append(z)
   z.set(qn('w:val'),'24')

def read_data(doc):
 ts=doc.tables
 a=np.array([[int(c.text) for c in r.cells[1:]] for r in ts[34].rows[1:]],dtype=np.int64)
 b=np.array([[int(c.text) for c in r.cells[1:]] for r in ts[36].rows[1:]],dtype=np.int64)
 dc=np.array([[int(c.text) for c in r.cells[1:]] for r in ts[37].rows[1:]],dtype=np.int64)
 l=[]
 for t in ts[38:43]:l.append(np.array([[float(c.text.replace('%','')) for c in r.cells[1:]] for r in t.rows[1:]]))
 res=[]
 for k,z in enumerate(l):
  assert np.array_equal(z[:,0],b[:,k]);assert np.array_equal(z[:,1],dc[:,k])
  hw=z[:,0]/z[:,5];sel=z[:,5]/z[:,2];comb=z[:,0]/z[:,2]
  assert np.allclose(hw*sel,comb,rtol=1e-13,atol=0)
  res.append({'cores':k+1,'fixed_hit_mean':float(z[:,6].mean()),'selected_hit_mean':float(z[:,4].mean()),'geom_hw':float(np.exp(np.log(hw).mean())),'geom_sel':float(np.exp(np.log(sel).mean())),'geom_comb':float(np.exp(np.log(comb).mean())),'hw_mean':float(hw.mean()),'sel_mean':float(sel.mean()),'combined_mean':float(comb.mean()),'hw_median':float(np.median(hw)),'spearman':float(spearmanr(z[:,6],hw).statistic),'faster':int((hw>1).sum()),'equal':int((hw==1).sum()),'slower':int((hw<1).sum())})
 ex=[]
 for sc,x in [('A',a[:,1:]),('B',b[:,1:]),('L2',np.column_stack([z[:,2] for z in l[1:]]))]:
  for j in range(3):
   ratio=x[:,j+1]/x[:,j];bad=np.where(ratio>1)[0]
   ex.append({'scene':sc,'from':j+2,'to':j+3,'faster':int((ratio<1).sum()),'equal':int((ratio==1).sum()),'slower':len(bad),'worst_pct':float(100*(ratio.max()-1)),'worst_case':f'case_{ratio.argmax()+1:03d}','bad_cases':[f'case_{i+1:03d}' for i in bad]})
 signs=[]
 for r in ts[30].rows[1:]:
  label=r.cells[0].text;nplus,zero,nminus=map(int,r.cells[1].text.split('/'));n=nplus+nminus
  p=min(1.,2*sum(math.comb(n,j) for j in range(min(nplus,nminus)+1))/2**n)
  signs.append({'group':label,'p':p,'p_bonferroni':min(1.,14*p)})
 original_tables=[[[c.text for c in r.cells] for r in t.rows] for t in ts[34:43]]
 return {'l2':res,'expansion':ex,'sign':signs,'source_data_rows':900},original_tables

EQ={
 '4-1':r'P=(g,a,\pi),\quad g:U\to S,\quad a:S\to\{0,\ldots,k-1\}',
 '4-2':r'\bigcup_{s\in S}g^{-1}(s)=U,\qquad g^{-1}(s)\cap g^{-1}(s^\prime)=\varnothing\ (s\ne s^\prime)',
 '4-3':r'\begin{aligned}E_{dep}&=\{(s,s^\prime):s\ne s^\prime,\ \exists u\in g^{-1}(s),u^\prime\in g^{-1}(s^\prime),u\prec u^\prime\},\\E_{ord}&=\bigcup_c\{(\pi_c(j),\pi_c(j+1))\}.\end{aligned}',
 '4-20':r'\mathop{\mathrm{lex\ min}}_{P\in\mathcal F_X}\left(T_X(P),D_{added}^{X}(P)\right),\qquad X\in\{A,B,L\}',
 '5-1':r'\mathop{\mathrm{lex\ min}}_{P\in\mathcal F_A}\left(T_A(P),D_{added}^{A}(P)\right)',
 '6-1':r'\mathop{\mathrm{lex\ min}}_{P\in\mathcal F_B}\left(T_B(P),D_{added}^{B}(P)\right)',
 '7-1':r'\mathop{\mathrm{lex\ min}}_{P\in\mathcal F_L}\left(T_L(P),D_{added}^{B}(P)\right)',
 '6-3':r'\mathcal C_{family}=\{(m,target):1\le m\le k,\quad target\in\{0,512,2048\}\}',
 '5-8':r'\begin{aligned}span+16\,\mathrm{ulp}(\cdot)&<need\quad\mathrm{or}\\rightmost+16\,\mathrm{ulp}(\cdot)&<ready+need,\quad need=duration+gap.\end{aligned}',
 '7-3':r'h_i=\frac{B_{hit,i}}{B_{hit,i}+B_{miss,i}}\quad(B_{hit,i}+B_{miss,i}>0)',
 '8-3':r'p=\min\left\{1,\ 2\sum_{j=0}^{\min(n_+,n_-)}\binom{n_++n_-}{j}2^{-(n_++n_-)}\right\}',
}

def additions(stats):
 A={243:[
 ('h2','4.9 下界的适用域与可核验判定'),
 ('h3','4.9.1 同一个不等式为何不能承担所有职责'),
 ('p',r'全局下界在固定图与核数下对全部合法方案成立，因此可以约束真正最优值。固定方案下界则只约束给定划分和核心归属的方案；改变归属后，各核负载和必需搬运都可能变化。将后者当作全局最优值的下界，会把“这个候选的资源很紧”误写成“所有方案都没有改进空间”。'),
 ('p',r'设当前成功方案的时间为 $T^*$。对于待评候选 $P$，若已证明 $L(P)\le T(P)$ 且 $L(P)>T^*$，则该候选不可能在主目标上胜出，可以跳过。若 $L(P)=T^*$，仍不能淘汰：它可能主目标持平而新增 COPY 更少。对活动核数至多为 $m$ 的整个候选族，则须使用对该族全部方案成立的 $LB_m$；其用途是类别淘汰，而不是仅对其中某一份方案的判断。'),
 ('h3','4.9.2 固定方案的乐观带宽界'),
 ('p',r'设固定方案在 Spill 插入前的计划读入、写出量为 $R$、$W$。不含 L2 时，全部这些工作均要经过共享 DDR，忽略释放时间和其他竞争只能使所需时间减少，因此总传输工作量给出下界。L2 场景中，所有写出仍必须经过 DDR；将所有读取乐观地视为可以自由分配到两个池，并忽略首次未命中、FIFO 淘汰及容量，可以得到更弱但安全的资源松弛。'),
 ('eq',('4-24',r'L_{copy}^{A/B}(P)=\left\lceil\frac{R+W}{60}\right\rceil')),
 ('eq',('4-25',r'L_{copy}^{L}(P)=\max\left\{\left\lceil\frac{W}{60}\right\rceil,\left\lceil\frac{R+W}{60+250}\right\rceil\right\}')),
 ('p',r'第一项来自写出只能使用 DDR；第二项来自两个独立池在任意持续时间内能够提供的总服务量上限。真实执行还要满足读池资格、Pipe 占用、跨核释放和取整规则，因而该式不是对实际命中率的估计。最终固定方案界还取各核计算 Pipe 负载、计算关键路径与带宽界的最大值。'),
 ('h3','4.9.3 非空窗口的推导与示例'),
 ('p',r'对窗口集合 $U(\alpha,\beta,p)$ 中的任一操作，原依赖使它不能早于 $\alpha$ 开始；它完成后还至少有 $\beta$ 的计算必须顺序接续。因此，该集合内全部计算只能占据长度为 $T-\alpha-\beta$ 的区间。$k$ 条同类流水线在该区间最多提供 $k(T-\alpha-\beta)$ 的计算容量，将集合工作量与这一容量比较即得式(4-22)。'),
 ('p',r'作为代数示例，设两核下某个非空窗口的前缀和后缀阈值分别为10和8周期，窗口内同类计算总工作量为25周期，则任何可行执行至少需要 $10+8+\lceil25/2\rceil=31$ 周期。这里的25不是图中某一次真实测量；它只演示“局部必需工作量加两端依赖”的计算方法。若集合为空，该例的工作量约束并不存在，不能直接把10和8相加作为新下界。'),
 ('h3','4.9.4 条件证书与程序路径'),
 ('p',r'以 $LB>0$ 表示对全部方案成立的全局界，则当前解与真正最优值之间的相对差距满足'),
 ('eq',('4-26',r'0\le\frac{T(P^*)-T_{opt}}{T_{opt}}\le\frac{T(P^*)}{LB}-1.')),
 ('p',r'例如，假设一个实例的全局界为1000周期、当前解为1025周期，则可认证的差距上限为2.5%；若当前解为1500周期，则只能给出50%的上限。后一个界可能很松，不能据此判断实际解真的比最优值慢50%。这两个数值例子均为条件演示，不计入100张图的实验结果。'),
 ('p',r'代码中的近优门控只是跳过其所覆盖的控制阶段。外层插入和微批还有独立的触发与剩余预算检查，不能由证书成立推断程序已经退出。本稿不汇报未经逐请求日志确认的证书触发率；可核验对象是每次实际保存的界、成功方案和门控记录。'),
 ],
 667:[
 ('h2','7.7 冻结结果的补充统计与分解'),
 ('h3','7.7.1 两类命中率的逐核核对'),
 ('p',r'前文表7-3的命中率属于固定 B 方案开启 L2 的对照，不能直接写到 L2 独立选优方案名下。表7-11按附录的100个逐例百分比重新取算术平均，明确两种方案的区别；未进行新求解或新评分。命中率原表保留四位小数，复算也以该精度为基础。'),
 ('table',('表7-11 固定方案与独立选优方案的命中率（%）',['核数','固定 B 方案','L2 独立选优','固定方案命中率—配置比秩相关'],[[str(z['cores']),f"{z['fixed_hit_mean']:.4f}",f"{z['selected_hit_mean']:.4f}",f"{z['spearman']:.3f}"] for z in stats['l2']])),
 ('p',r'五核两类命中率分别为23.4614%和25.3143%。相关系数按同一图上的固定方案命中率与同方案配置比计算，不把两个方案的命中率混合。正相关表示在这批图中较高命中率通常伴随较大配置收益，不表示任意两张图都按命中率排序，也不排除开启缓存后个别方案更慢。'),
 ('h3','7.7.2 几何均值为何能够保留乘积关系'),
 ('p',r'令逐例配置比、选解比与综合比分别为 $x_i$、$y_i$、$z_i=x_iy_i$。由于这些比值均为正，先取对数再平均可以保持加法关系；指数还原后得到几何均值的乘积分解：'),
 ('eq',('7-12',r'G_x=\exp\left(\frac1n\sum_{i=1}^{n}\ln x_i\right),\quad G_y=\exp\left(\frac1n\sum_{i=1}^{n}\ln y_i\right)')),
 ('eq',('7-13',r'G_z=\exp\left(\frac1n\sum_i\ln(x_iy_i)\right)=G_xG_y.')),
 ('table',('表7-12 配置、选解与综合比的几何均值',['核数','配置效应','选解效应','综合比'],[[str(z['cores']),f"{z['geom_hw']:.6f}",f"{z['geom_sel']:.6f}",f"{z['geom_comb']:.6f}"] for z in stats['l2']])),
 ('p',r'五核三个几何均值约为1.016317、1.009934和1.026414，在未舍入值下严格满足乘积关系。几何均值是补充的乘性统计，不能替代题面要求的逐图加速比算术平均。正文主结果曲线、表10-1和逐例指标均保持题目口径。'),
 ('p',r'算术均值的情况不同。定义以 $n$ 为分母的经验协方差，则'),
 ('eq',('7-14',r'\overline{xy}=\bar x\bar y+\frac1n\sum_i(x_i-\bar x)(y_i-\bar y).')),
 ('p',r'只有协方差恰为零等特殊情形下，算术平均的综合比才等于两个算术平均效应的乘积。因此，直接用1.027400除以1.016766不能作为平均选解比的计算方法。该区别是聚合口径问题，不是评价器或实验数据产生矛盾。'),
 ('h3','7.7.3 比值与时间降幅不可互换'),
 ('p',r'对单例，时间降幅为 $d_i=1-1/x_i$，而比值超出1的部分为 $x_i-1$，两者在小幅变化时接近但并不相同。汇总时也有 $\bar d=1-\frac1n\sum_i1/x_i$，不能直接将平均配置比代入单例公式得到平均时间降幅。本文保留比值、百分比及其分母定义，避免把“约1.7%的比值超额”写成所有图普遍节省1.7%的时间。'),
 ],
 723:[
 ('h2','8.8 核数扩展的逐例非单调检验'),
 ('p',r'逐核平均曲线递增，不等于每张图增加一核都更快。利用附录中相邻核数的实际终点，按同一图计算 $T_{k+1}-T_k$，在2至5核的三个相邻转移上分别统计更快、持平和更慢。表8-6中每行仍为100张图，不把同一图跨核重复出现视为独立样本。'),
 ('table',('表8-6 相邻核数扩展的逐例结果（每行100图）',['场景','核数变化','更快/持平/更慢','最差增时/%','最差用例'],[[z['scene'],f"{z['from']}→{z['to']}",f"{z['faster']}/{z['equal']}/{z['slower']}",f"{z['worst_pct']:.4f}",z['worst_case'].replace('case_','')] for z in stats['expansion']])),
 ('p',r'最差增时以 $100(T_{k+1}/T_k-1)$ 计算，正值表示更慢。增加核心理论上允许保留原有分配并添加空队列，但这里比较的是两次独立有限搜索得到的最终方案，不是把原方案无条件嵌入后重新选择；因此观测非单调不与“允许空核心”相矛盾。'),
 ('p',r'这些记录与相对 C3 的退化是两类不同事件。前者比较本方法相邻核数的方案，后者比较同核数两种方法的方案。同一用例可以在一种比较中改善、在另一种比较中退化，不能相加作为统一失败率。case_062 的低核数对照问题以及 case_039 的 A 四核至五核变化，应分别沿各自比较轴诊断。'),
 ('p',r'本检验只复用归档周期，既没有为非单调实例重新搜索，也没有将低核数优胜方案追加到高核数候选集。若后续采用跨核保底机制，需要重新定义预算、运行并报告全部受影响结果；不能只把附录中的高核数数值替换为较小值。'),
 ],
 809:[
 ('h3','B.1 两种复现任务应分开'),
 ('p',r'第一种任务是复核论文数值：读取归档的方案，用同一版本的官方评价器与固定配置重新评价，再对照 Makespan、新增 COPY 和缓存命中率。第二种任务是重新运行搜索：从原图生成新候选并受墙钟预算限制。前者检查既有结果与方案是否一致，后者检查求解流程能否再次产生可行方案；两者不是同一项测试。'),
 ('h3','B.2 输入、环境与失败记录'),
 ('p',r'准备100张原始 case_*.json、固定 config.txt 及完整求解源码。原始图应来自题目材料，不以表格反推或合成替代。执行命令时工作目录切换到 reproduction/NPU_R12，使用全新或空输出目录，避免将另一请求的结果误作当前结果。运行期间保存解释器版本、平台、输入和配置哈希，以及各阶段实际状态。'),
 ('p',r'常规请求仍使用600秒，case_087 的 A2至A5分别保留首次失败和1800秒补跑目录。输出方案顶层仅保留题目规定的两个字段；日志、报告及哈希文件作为外部支撑材料保存。若缺少原始图、归档方案或程序版本，必须说明缺失，不能只凭同名文件认定复现成功。'),
 ('h3','B.3 表格与统计复算'),
 ('p',r'新增统计只读表 A-1至A-9：逐图比值先计算再平均，命中率按已有四位小数记录复算；按用例编号连接无 L2 与 L2 表，断言100行编号完整、唯一且一一对应。几何分解检验使用未舍入值核对乘积；相邻核数比较按原始整数周期计算。该流程验证的是表内算术与统计一致性，不替代完整事件模拟。'),
 ],
 }
 return A

class Rich:
 def __init__(self, strings):
  formulas=sorted(set(f for s in strings for f in re.findall(r'\$([^$]+)\$',s)))
  formulas+=sorted(set(EQ.values())-set(formulas));self.math={}
  if not formulas:return
  with tempfile.TemporaryDirectory() as td:
   md='\n\n'.join(f'FORMULA{i:04d}\n\n$${f}$$' for i,f in enumerate(formulas));out=Path(td)/'math.docx'
   pypandoc.convert_text(md,'docx',format='markdown',outputfile=str(out),extra_args=['--standalone'])
   d=Document(out);current=None
   for p in d.paragraphs:
    if p.text.startswith('FORMULA'):current=int(p.text[7:])
    ms=p._p.xpath('.//m:oMath')
    if ms and current is not None:self.math[formulas[current]]=copy.deepcopy(ms[0]);current=None
  missing=[f for f in formulas if f not in self.math]
  if missing:raise ValueError('Missing math: '+repr(missing))
 def fill(self,p,text):
  for el in list(p._p):
   if el.tag!=qn('w:pPr'):p._p.remove(el)
  for j,part in enumerate(re.split(r'\$([^$]+)\$',text)):
   if j%2:p._p.append(copy.deepcopy(self.math[part]))
   else:p.add_run(part)

def borders(t,visible=True):
 pr=t._tbl.tblPr
 for old in list(pr.findall(qn('w:tblBorders'))):pr.remove(old)
 b=node('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  b.append(node('w:'+side,**{'w:val':('single' if visible and side in ('top','bottom') else 'nil'),'w:sz':'8','w:color':'000000'}))
 pr.append(b)

def widths(t,cm):
 t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 grid=t._tbl.tblGrid
 for x in list(grid):grid.remove(x)
 for w in cm:grid.append(node('w:gridCol',**{'w:w':int(Cm(w).twips)}))
 for r in t.rows:
  for c,w in zip(r.cells,cm):c.width=Cm(w)
 pr=t._tbl.tblPr
 for tag in ['w:tblInd','w:tblCellSpacing']:
  for x in list(pr.findall(qn(tag))):pr.remove(x)
 tw=pr.find(qn('w:tblW'));tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),str(int(Cm(sum(cm)).twips)))

def style_table(t,cm=None,eq=False):
 if cm is None:cm=[WIDTH/len(t.columns)]*len(t.columns)
 widths(t,cm);borders(t,not eq)
 pr=t._tbl.tblPr
 for old in list(pr.findall(qn('w:tblCellMar'))):pr.remove(old)
 mar=node('w:tblCellMar')
 for side in ('top','bottom','left','right'):mar.append(node('w:'+side,**{'w:w':('10' if side in ('top','bottom') else '45'),'w:type':'dxa'}))
 pr.append(mar)
 for ri,r in enumerate(t.rows):
  rp=r._tr.get_or_add_trPr()
  for tag in ['w:trHeight','w:tblHeader','w:cantSplit']:
   for old in list(rp.findall(qn(tag))):rp.remove(old)
  rp.append(node('w:cantSplit'))
  if ri==0 and not eq:rp.append(node('w:tblHeader'))
  for ci,c in enumerate(r.cells):
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   cp=c._tc.get_or_add_tcPr()
   for tag in ['w:tcBorders','w:shd','w:noWrap']:
    for old in list(cp.findall(qn(tag))):cp.remove(old)
   if ri==0 and not eq:
    b=node('w:tcBorders');b.append(node('w:bottom',**{'w:val':'single','w:sz':'4','w:color':'000000'}));cp.append(b)
   for p in c.paragraphs:
    para_format(p,'eq' if eq else 'table')
    for run in p.runs:set_font(run,12,ri==0 and not eq)
    if eq and ci==1:p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    if not eq and len(c.text)>30:p.alignment=WD_ALIGN_PARAGRAPH.LEFT

def build(source,outdir):
 raw=source.read_bytes();assert hashlib.sha256(raw).hexdigest()==EXPECTED,'Source Word differs; reconcile rather than overwrite.'
 d=Document(source);ps=list(d.paragraphs);ts=list(d.tables)
 stats,archive=read_data(d);adds=additions(stats)
 edits=json.loads((HERE/'edits.json').read_text());edits={int(k):{'text':v} if isinstance(v,str) else v for k,v in edits.items()}
 strings=[v['text'] for v in edits.values()]
 for blocks in adds.values():
  for kind,data in blocks:
   if kind=='eq':EQ[data[0]]=data[1]
   elif kind!='table':strings.append(data)
 rich=Rich(strings)
 for idx,record in edits.items():
  assert 'expected' not in record or ps[idx].text==record['expected'],f'Paragraph drift: {idx}'
  rich.fill(ps[idx],record['text'])
 first=ps[10]._p
 while d.element.body[0] is not first:d.element.body.remove(d.element.body[0])
 for i in (11,):
  if ps[i]._p.getparent() is not None:ps[i]._p.getparent().remove(ps[i]._p)
 for pr in d.element.body.xpath('.//w:pPr/w:sectPr'):pr.getparent().remove(pr)
 for sect in d.sections:
  sect.page_width=Cm(21);sect.page_height=Cm(29.7);sect.left_margin=Cm(2.25);sect.right_margin=Cm(2.25);sect.top_margin=Cm(2.5);sect.bottom_margin=Cm(2.0);sect.header_distance=Cm(.8);sect.footer_distance=Cm(1.1)
  sect.different_first_page_header_footer=False
  for x in list(sect._sectPr.findall(qn('w:pgNumType'))):sect._sectPr.remove(x)
  sect._sectPr.append(node('w:pgNumType',**{'w:start':'1'}))
  for h in [sect.header,sect.first_page_header,sect.even_page_header]:
   for p in h.paragraphs:p.clear()
  for f in [sect.footer,sect.first_page_footer,sect.even_page_footer]:
   for p in f.paragraphs:p.clear()
  p=sect.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=1
  run=p.add_run();run._r.append(node('w:fldChar',**{'w:fldCharType':'begin'}));r=p.add_run();it=node('w:instrText');it.text=' PAGE ';r._r.append(it)
  p.add_run()._r.append(node('w:fldChar',**{'w:fldCharType':'end'}))
  for r in p.runs:set_font(r,12)
 for style in d.styles:
  if style.type in (1,2):
   style.font.name='Times New Roman';style.font.size=Pt(12)
   rp=style.element.get_or_add_rPr();rf=rp.find(qn('w:rFonts'))
   if rf is None:rf=node('w:rFonts');rp.insert(0,rf)
   rf.set(qn('w:eastAsia'),'宋体');rf.set(qn('w:ascii'),'Times New Roman');rf.set(qn('w:hAnsi'),'Times New Roman')
  if style.type==1:style.paragraph_format.line_spacing=1.0
 for p in d.paragraphs:
  sn=p.style.name
  kind='h'+sn[-1] if sn.startswith('Heading ') else ('caption' if sn in ('图题','表题') else ('algo' if sn=='Block Text' else ('code' if sn=='Source Code' else 'body')))
  para_format(p,kind)
  if sn=='表题':p.paragraph_format.keep_with_next=True
  if p._p.xpath('.//w:drawing'):
   p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(1);p.paragraph_format.keep_with_next=True
  if p.text.startswith('['):p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(5)
 for i in [20,765,786,802]:ps[i].paragraph_format.page_break_before=True
 para_format(ps[10],'h1')
 for r in ps[10].runs:set_font(r,16,True,'黑体')
 ps[10].paragraph_format.space_before=Pt(2);ps[10].paragraph_format.space_after=Pt(10)
 ps[12].alignment=WD_ALIGN_PARAGRAPH.CENTER;ps[12].paragraph_format.first_line_indent=Pt(0)
 for r in ps[12].runs:set_font(r,12,True)
 ts[2].cell(3,0).text='最大分层宽度（操作数）'
 ts[5].cell(1,1).text='完整等价编码需另外实现与验证';ts[5].cell(1,2).text='取决于规模及具体编码';ts[5].cell(2,1).text='代理需校准；也可直接接官方评价';ts[5].cell(2,2).text='取决于评价成本与搜索预算';ts[5].cell(3,4).text='候选确定；预算截断可能改变终点'
 ts[7].cell(2,3).text='仅跳过受门控控制阶段；外层插入与微批仍可运行'
 ts[8].cell(4,2).text='固定散列不自动保证随机模型的期望深度'
 ts[8].cell(5,1).text='计算界线性；窗口维护 O(n log n)'
 ts[18].cell(5,0).text='固定 B 方案命中率均值';ts[18].cell(5,2).text='逐图字节命中率的算术平均，非独立选优方案'
 ts[20].cell(0,5).text='固定 B 方案\n命中率均值'
 ts[26].cell(10,3).text='○';ts[26].cell(10,2).text='确有 Spill 时诊断并作窗口重排，A 不启用'
 ts[28].cell(4,3).text='符号检验原始 p<10⁻²²；14组校正后最大约1.226×10⁻²¹'
 ts[28].cell(5,3).text='归档结果均不慢于 REF；描述结构与性能的秩相关'
 for ti in range(34,43):
  t=ts[ti]
  for r in t.rows[1:]:r.cells[0].text=r.cells[0].text.replace('case_','')
  if ti>=38:
   for r in t.rows:
    cells=list(r._tr.tc_lst)
    r._tr.remove(cells[2]);r._tr.remove(cells[1])
   grid=t._tbl.tblGrid
   for x in [grid[2],grid[1]]:grid.remove(x)
   heads=['用例','L2选优\n时间','L2选优\n新增字节','L2选优\n命中率/%','固定方案\nL2时间','固定方案\n命中率/%']
   for c,h in zip(t.rows[0].cells,heads):c.text=h
  else:
   if ti==34:ts[ti].cell(0,1).text='REF单核'
 layouts={1:[5.1,2.15,2.2,2.2,2.2,2.65],2:[5.4,1.9,2.2,2.2,2.2,2.6],3:[4.4,8.4,3.7],4:[6.2,4.6,5.7],5:[2.7,3.9,3.5,3.4,3.0],6:[2.6,4.4,4.75,4.75],7:[3.0,4.1,4.2,5.2],8:[4.3,5.2,7.0],9:[5.5,2.75,2.75,2.75,2.75],10:[1.1,3.2,3.5,4.1,4.6],11:[.8,2.7,2.7,1.2,2.0,3.0,4.1],13:[2.1,1.7,3.2,2.5,3.4,3.6],14:[5.0,2.3,2.3,2.3,2.3,2.3],15:[.8,2.5,2.9,3.6,3.3,3.4],16:[3.8,8.9,3.8],18:[5.4,4.2,6.9],19:[5.0,2.3,2.3,2.3,2.3,2.3],20:[.85,2.65,2.65,2.65,3.4,2.75,1.55],21:[3.1,1.4,5.6,6.4],22:[3.9,4.2,4.2,4.2],24:[1.1,3.2,3.5,4.1,4.6],25:[1.5,5.9,4.7,4.4],26:[.65,3.3,5.4,.7,.7,.7,5.05],27:[1.3,3.8,3.8,3.8,3.8],28:[2.4,4.5,3.0,3.6,3.0],29:[2.2,2.4,2.1,3.6,3.6,2.6],30:[2.8,4.3,3.3,3.0,3.1],31:[2.6,4.3,5.3,4.3],32:[6.8,3.25,3.25,3.2],33:[.8,2.5,2.5,2.5,2.5,2.8,2.9],34:[1.4,3.05,3.0,3.0,3.0,3.05],35:[1.5,3.75,3.75,3.75,3.75],36:[1.4,3.05,3.0,3.0,3.0,3.05],37:[1.4,3.05,3.0,3.0,3.0,3.05],43:[3.7,6.0,6.8]}
 for i,t in enumerate(ts[1:],1):
  cm=layouts.get(i)
  if i>=38 and i<=42:cm=[1.3,3.0,3.0,2.9,3.3,3.0]
  if cm:cm=[x*WIDTH/sum(cm) for x in cm]
  style_table(t,cm)
 for shape in d.inline_shapes:
  ow,oh=shape.width,shape.height;w=min(Cm(15.5),ow);h=int(w*oh/ow)
  if h>Cm(13.3):h=Cm(13.3);w=int(h*ow/oh)
  shape.width=int(w);shape.height=int(h)
 eq_count=0
 for p in list(d.paragraphs):
  match=re.fullmatch(r'\s*\((\d+-\d+)\)\s*',p.text)
  if not match or not p._p.xpath('.//m:oMath'):continue
  number=match[1];table=d.add_table(rows=1,cols=2);p._p.addprevious(table._tbl)
  dest=table.cell(0,0).paragraphs[0]
  if number in EQ:dest._p.append(copy.deepcopy(rich.math[EQ[number]]))
  else:
   for mo in p._p.xpath('.//m:oMath'):dest._p.append(copy.deepcopy(mo))
  table.cell(0,1).text='('+number+')';style_table(table,[15.2,1.3],True)
  p._p.getparent().remove(p._p);eq_count+=1
 def add_p(anchor,text,kind):
  p=Paragraph(OxmlElement('w:p'),d._body);anchor.addprevious(p._p);p.style=d.styles['Heading '+kind[-1]] if kind.startswith('h') else d.styles['Body Text'];rich.fill(p,text);para_format(p,kind if kind!='p' else 'body');return p
 for idx,blocks in adds.items():
  anchor=ps[idx]._p
  for kind,data in blocks:
   if kind=='eq':
    number,formula=data;t=d.add_table(rows=1,cols=2);anchor.addprevious(t._tbl);t.cell(0,0).paragraphs[0]._p.append(copy.deepcopy(rich.math[formula]));t.cell(0,1).text='('+number+')';style_table(t,[15.2,1.3],True);eq_count+=1
   elif kind=='table':
    caption,heads,rows=data;p=add_p(anchor,caption,'caption');p.paragraph_format.keep_with_next=True
    t=d.add_table(rows=1,cols=len(heads));anchor.addprevious(t._tbl)
    for c,h in zip(t.rows[0].cells,heads):c.text=h
    for row in rows:
     for c,value in zip(t.add_row().cells,row):c.text=str(value)
    style_table(t)
   else:add_p(anchor,data,kind)
 for idx,comment in {197:'可行域由覆盖、唯一排布、联合无环和完整官方执行成功共同定义；必要的时序、容量、带宽与计账约束分别见式(4-2)至式(4-19)。',251:'场景 A：每个子图独立成为 Task，释放规则见式(4-11)至式(4-13)，边界计账见式(4-18)。',379:'场景 B：同核子图合并成一个 Task，跨核同步见式(4-14)，边界计账见式(4-19)。'}.items():
  add_p(ps[idx+1]._p,comment,'note')
 d.core_properties.author='';d.core_properties.last_modified_by='';d.core_properties.comments='';d.core_properties.title='基于可证下界与结构化搜索的多核 NPU 切图调度';d.core_properties.subject='完整增强核验稿'
 for x in list(d.element.body.xpath('.//w:bookmarkStart | .//w:bookmarkEnd')):x.getparent().remove(x)
 rows_kept=0
 for off,t in enumerate(ts[34:43]):
  for ri,r in enumerate(t.rows[1:],1):
   got=[c.text for c in r.cells];orig=archive[off][ri];orig[0]=orig[0].replace('case_','')
   if off>=4:orig=[orig[j] for j in (0,3,4,5,6,7)]
   assert got==orig,(off,ri,got,orig)
   rows_kept+=1
 assert rows_kept==900
 outdir.mkdir(parents=True,exist_ok=True);dest=outdir/(NAME+'.docx');d.save(dest)
 with zipfile.ZipFile(dest) as zin:
  entries={n:zin.read(n) for n in zin.namelist()}
 for n,b in list(entries.items()):
  if n.startswith('word/media/') and n.lower().endswith('.png'):
   im=Image.open(io.BytesIO(b));im.thumbnail((2400,2400));buf=io.BytesIO();im.save(buf,format='PNG',optimize=True);entries[n]=buf.getvalue()
 with zipfile.ZipFile(dest,'w',zipfile.ZIP_DEFLATED) as zout:
  for n,b in entries.items():zout.writestr(n,b)
 report={'source_sha256':EXPECTED,'patched_paragraphs':len(edits),'preserved_original_figures':len(d.inline_shapes),'numbered_equations':eq_count,'archival_rows_verified':rows_kept,'derived_statistics':stats,'status':'built; PDF page/visual validation required','fonts':'DOCX specifies SimSun/SimHei; Linux PDF uses Noto CJK fallback. Formal font export requires team review.'}
 (outdir/'校验记录.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
 (outdir/'补充统计.json').write_text(json.dumps(stats,ensure_ascii=False,indent=2))
 print(json.dumps({k:v for k,v in report.items() if k!='derived_statistics'},ensure_ascii=False,indent=2));print(dest)
 return dest

if __name__=='__main__':
 ap=argparse.ArgumentParser();ap.add_argument('--source',type=Path,default=HERE.parent/'paper_full_v1/out/R12论文_v1.docx');ap.add_argument('--out',type=Path,default=HERE/'out');args=ap.parse_args();build(args.source,args.out)
