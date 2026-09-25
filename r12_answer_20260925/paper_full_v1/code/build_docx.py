"""把合稿中间件转成Word草稿：三线表、公式右对齐编号、图题表题、正文首行缩进。

依赖：pypandoc-binary（pandoc 3.x）、python-docx、Pillow。输出 ../out/R12论文_v1.docx。
figures_v3 图片顶部带内部图号标题，这里裁掉标题带后再嵌入（原图不改）。
"""
from __future__ import annotations

import copy
import re
import subprocess
from pathlib import Path

import pypandoc
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
OUT = ROOT / "out"
FIGDIR = OUT / "fig"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
TITLE = "通用神经网络处理器下的多核调度问题"


def crop_title_band(src: Path, dst: Path) -> None:
    """裁掉图片顶部的标题行：找到第一段非白行及其后的空白间隔，从间隔处开始保留。"""
    im = Image.open(src).convert("RGB")
    w, h = im.size
    gray = im.convert("L")
    px = gray.load()
    step = max(1, w // 400)
    def dark(y):
        return any(px[x, y] < 200 for x in range(0, w, step))
    y = 0
    while y < h * 0.12 and not dark(y):
        y += 1
    top_text = y
    while y < h * 0.15 and dark(y):
        y += 1
    if top_text < h * 0.12 and y < h * 0.15:
        gap_start = y
        while y < h * 0.2 and not dark(y):
            y += 1
        cut = max(gap_start, y - int(0.01 * h))
        im = im.crop((0, cut, w, h))
    dst.parent.mkdir(parents=True, exist_ok=True)
    im.save(dst, dpi=(600, 600))


def prepare_markdown() -> Path:
    text = (OUT / "_pandoc_input.md").read_text(encoding="utf-8")
    def repl(m):
        src = Path(m.group(1))
        dst = FIGDIR / src.name
        if "figures_v3" in src.as_posix():
            crop_title_band(src, dst)
        else:
            Image.open(src).save(dst, dpi=(600, 600))
        return f"![]({dst.as_posix()}){{width=15.5cm}}"
    text = re.sub(r"!\[\]\(([^)]+)\)\{width=15\.5cm\}", repl, text)
    # Word中多字母\mathrm按字母拆分，统一改为\text以保持正体整词
    text = re.sub(r"\\mathrm\{([A-Za-z][A-Za-z ]+)\}", r"\\text{\1}", text)
    text = text.replace("{width=15.5cm}", "{width=14cm}")
    text = f"::: {{custom-style=\"论文题目\"}}\n{TITLE}\n:::\n\n" + text
    path = OUT / "_pandoc_ready.md"
    path.write_text(text, encoding="utf-8")
    return path


def set_run_fonts(style, east="宋体", west="Times New Roman", size=None, bold=None):
    style.font.name = west
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(attr), west)
    fonts.set(qn("w:eastAsia"), east)
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold


def reference_doc() -> Path:
    ref = OUT / "_reference.docx"
    data = subprocess.run([pypandoc.get_pandoc_path(), "--print-default-data-file", "reference.docx"],
                          capture_output=True, check=True).stdout
    ref.write_bytes(data)
    doc = Document(ref)
    st = doc.styles
    for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text"):
        if name in [s.name for s in st]:
            set_run_fonts(st[name], size=12)
            pf = st[name].paragraph_format
            pf.line_spacing = 1.25; pf.space_before = Pt(0); pf.space_after = Pt(0)
    for lvl, size in ((1, 16), (2, 14), (3, 12), (4, 12)):
        s = st[f"Heading {lvl}"]
        set_run_fonts(s, east="黑体", west="Times New Roman", size=size, bold=True)
        s.font.color.rgb = None
        s.paragraph_format.space_before = Pt(12 if lvl == 1 else 6)
        s.paragraph_format.space_after = Pt(6 if lvl == 1 else 3)
        if lvl == 1:
            s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for name in ("Title", "Subtitle"):
        if name in [s.name for s in st]:
            set_run_fonts(st[name], east="黑体", size=18, bold=True)
    doc.save(ref)
    return ref


def number_equations(doc) -> int:
    body = doc.element.body
    n = 0
    for p in list(body.iter(qn("w:p"))):
        text = "".join(t.text or "" for t in p.iter(qn("w:t")))
        m = re.fullmatch(r"EQNO\(([^)]+)\)", text.strip())
        if not m:
            continue
        prev = p.getprevious()
        para = prev.find(f"{{{M_NS}}}oMathPara") if prev is not None else None
        if prev is None or para is None:
            raise SystemExit(f"公式编号 {m.group(1)} 前没有展示公式")
        omath = para.find(f"{{{M_NS}}}oMath")
        prev.remove(para)
        pf = Paragraph(prev, doc._body).paragraph_format
        pf.tab_stops.add_tab_stop(Cm(8.0), WD_TAB_ALIGNMENT.CENTER)
        pf.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT)
        pf.first_line_indent = Cm(0)
        def run_tab():
            r = OxmlElement("w:r"); r.append(OxmlElement("w:tab")); return r
        prev.append(run_tab())
        prev.append(omath)
        r = run_tab()
        t = OxmlElement("w:t"); t.text = f"({m.group(1)})"; r.append(t)
        prev.append(r)
        p.getparent().remove(p)
        n += 1
    return n


def border(tag, sz):
    b = OxmlElement(f"w:{tag}")
    b.set(qn("w:val"), "single" if sz else "nil"); b.set(qn("w:sz"), str(sz)); b.set(qn("w:space"), "0"); b.set(qn("w:color"), "000000")
    return b


def three_line(table, small=False) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblpr = table._tbl.tblPr
    for old in tblpr.findall(qn("w:tblBorders")):
        tblpr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for tag, sz in (("top", 12), ("left", 0), ("bottom", 12), ("right", 0), ("insideH", 0), ("insideV", 0)):
        borders.append(border(tag, sz))
    tblpr.append(borders)
    jc = tblpr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc"); tblpr.append(jc)
    jc.set(qn("w:val"), "center")
    header = table.rows[0]
    for cell in header.cells:
        tcpr = cell._tc.get_or_add_tcPr()
        for old in tcpr.findall(qn("w:tcBorders")):
            tcpr.remove(old)
        tb = OxmlElement("w:tcBorders"); tb.append(border("bottom", 4)); tcpr.append(tb)
    size = Pt(8 if small else 10.5)
    for row in table.rows:
        for cell in row.cells:
            for par in cell.paragraphs:
                par.paragraph_format.first_line_indent = Cm(0)
                par.paragraph_format.line_spacing = 1.0
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in par.runs:
                    run.font.size = size
    for cell in header.cells:
        for par in cell.paragraphs:
            for run in par.runs:
                run.font.bold = True


def style_paragraphs(doc) -> None:
    for par in doc.paragraphs:
        name = par.style.name
        txt = par.text.strip()
        if name in ("表题", "图题"):
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.first_line_indent = Cm(0)
            par.paragraph_format.space_before = Pt(3 if name == "表题" else 0)
            par.paragraph_format.space_after = Pt(3 if name == "表题" else 6)
            for run in par.runs:
                run.font.size = Pt(10.5); run.font.bold = True
                rf = run._element.get_or_add_rPr().get_or_add_rFonts()
                rf.set(qn("w:eastAsia"), "黑体")
            continue
        if name == "论文题目":
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in par.runs:
                run.font.size = Pt(18); run.font.bold = True
            continue
        has_img = bool(par._element.findall(".//" + qn("w:drawing")))
        if has_img:
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.first_line_indent = Cm(0)
            continue
        if name.startswith("Heading"):
            if txt in ("1 问题重述", "参考文献", "附录"):
                par.paragraph_format.page_break_before = True
            continue
        if not txt:
            continue
        if par._element.find(f".//{{{M_NS}}}oMath") is not None and par._element.find(qn("w:pPr") + "/" + qn("w:tabs")) is not None:
            continue
        if re.match(r"^\[\d+\] ", txt):          # 参考文献条目：悬挂缩进
            par.paragraph_format.left_indent = Cm(0.9)
            par.paragraph_format.first_line_indent = Cm(-0.9)
            continue
        if name in ("Body Text", "First Paragraph", "Normal"):
            par.paragraph_format.first_line_indent = Cm(0.85)


ORDER = {
    "pPr": ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr", "suppressLineNumbers",
            "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
            "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
            "suppressOverlap", "jc", "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle",
            "rPr", "sectPr", "pPrChange"],
    "rPr": ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline", "shadow", "emboss",
            "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing", "w", "kern", "position", "sz", "szCs",
            "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout",
            "specVanish", "oMath"],
    "tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc",
              "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription"],
    "tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "textDirection",
             "tcFitText", "vAlign", "hideMark"],
}


def reorder_children(doc) -> None:
    """按OOXML架构顺序重排属性节点的子元素，避免Word报“文件已损坏”。"""
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    roots = [doc.element, doc.styles.element]
    for root in roots:
        for tag, order in ORDER.items():
            rank = {W + name: i for i, name in enumerate(order)}
            for el in root.iter(W + tag):
                kids = list(el)
                known = [k for k in kids if k.tag in rank]
                if [k.tag for k in known] == sorted((k.tag for k in known), key=rank.get):
                    continue
                for k in kids:
                    el.remove(k)
                for k in sorted(kids, key=lambda k: rank.get(k.tag, len(order))):
                    el.append(k)


def main() -> None:
    FIGDIR.mkdir(parents=True, exist_ok=True)
    md = prepare_markdown()
    ref = reference_doc()
    out = OUT / "R12论文_v1.docx"
    pypandoc.convert_file(str(md), "docx", outputfile=str(out), format="markdown+pipe_tables+tex_math_dollars+fenced_divs+link_attributes",
                          extra_args=[f"--reference-doc={ref}", "--resource-path", str(ROOT)])
    doc = Document(out)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.5))
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    n = number_equations(doc)
    in_appendix = False
    body_items = list(doc.element.body.iterchildren())
    appendix_start = None
    for el in body_items:
        if el.tag == qn("w:p"):
            t = "".join(x.text or "" for x in el.iter(qn("w:t"))).strip()
            if t == "附录":
                appendix_start = el
    seen_appendix = False
    for el in body_items:
        if el is appendix_start:
            seen_appendix = True
        if el.tag == qn("w:tbl"):
            from docx.table import Table
            three_line(Table(el, doc), small=seen_appendix)
    style_paragraphs(doc)
    reorder_children(doc)
    doc.save(out)
    print(f"docx -> {out}  (numbered equations: {n}, tables: {len(doc.tables)})")


if __name__ == "__main__":
    main()
