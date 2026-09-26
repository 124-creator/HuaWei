# 本程序及代码是在人工智能工具辅助下完成的。
# 人工智能工具：Claude Code（Anthropic）；版本/型号与发布日期见论文附录D表D-1，由参赛队核实。
# 本脚本只做提交前整理，不在支撑材料包内；不运行求解器或官方评估器，不改任何实验数值。
"""生成正式提交文件：
1. 论文：以LBSS提交版Word为底稿（校验SHA256），加官方模板封面（第0页，学校/参赛队号/队员姓名），
   并入《细节核查记录_20260926.md》列明的标点与用字润色；页码自摘要页起为1。
2. 支撑材料：调用LBSS轮的 support.package_support 重建“支撑材料/”与zip（结构与论文附录B/C一致）。
用法：python3 tools/prepare_submission.py [--team-no 参赛队号]
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import io
import json
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

HERE = Path(__file__).resolve().parents[1]                     # 提交准备_2026A题/
REPO = HERE.parent
R12 = REPO / "r12_answer_20260925"
SUB = R12 / "paper_full_v2_expanded_20260926/submission_20260926"
SRC_DOCX = SUB / "out/多核NPU切图调度论文.docx"
SRC_SHA = "731421f1e173b28e263be5a167c350d0bb779f3bc4b9f90b63332891821e1558"
TEMPLATE = R12 / "paper_full_v1/template/附件3_论文模板_转docx.docx"
SCHOOL = "郑州航空工业管理学院"
MEMBERS = ["田中斐", "安镕基", "王玥"]
OUT_DOCX = HERE / "01_论文/多核NPU切图调度论文_提交版.docx"

sys.path.insert(0, str(R12 / "paper_full_v1/code"))
from build_docx import normalize_ooxml, el_text, reorder_children  # noqa: E402

# 《细节核查记录_20260926.md》：总起句句号改冒号；“只计”改“只计入”；“计0”改“按 0 计”。（原文, 新文, 应出现次数）
EDITS = [
    ("规模数据给出三条直接影响方法设计的信息。", "规模数据给出三条直接影响方法设计的信息：", 1),
    ("依赖结构又给出两条信息。", "依赖结构又给出两条信息：", 1),
    ("本文选择结构化候选与官方评价组合，主要考虑如下。", "本文选择结构化候选与官方评价组合，主要考虑如下：", 1),
    ("本章自身的结论可归为三点。", "本章自身的结论可归为三点：", 1),
    ("可以得到四条不依赖具体数字的认识。", "可以得到四条不依赖具体数字的认识：", 1),
    ("本文的方法特点概括如下。", "本文的方法特点概括如下：", 1),
    ("COPY 操作计0。", "COPY 操作按 0 计。", 1),
    ("只计原始计算周期", "只计入原始计算周期", 3),
    ("只计计算周期", "只计入计算周期", 2),
]


def sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def replace_in_paragraph(p, old: str, new: str) -> int:
    """在段落的w:t文字上替换（可跨run）；公式m:t不参与。返回替换次数。"""
    n = 0
    while True:
        ts = [t for t in p.iter(qn("w:t"))]
        text = "".join(t.text or "" for t in ts)
        i = text.find(old)
        if i < 0:
            return n
        pos, start_done = 0, False
        for t in ts:
            s = t.text or ""
            a, b = pos, pos + len(s)
            lo, hi = max(a, i), min(b, i + len(old))
            if lo < hi:
                if not start_done:
                    t.text = s[: lo - a] + new + s[hi - a:]
                    start_done = True
                else:
                    t.text = s[: lo - a] + s[hi - a:]
                t.set(qn("xml:space"), "preserve")
            pos = b
        n += 1


def apply_edits(doc) -> dict:
    counts = {}
    for old, new, want in EDITS:
        got = sum(replace_in_paragraph(p, old, new) for p in doc.element.body.iter(qn("w:p")))
        assert got == want, f"{old}: 期望{want}处，实际{got}处"
        counts[old] = got
    return counts


def fill_cover(table, team_no: str) -> None:
    cells = [tr.findall(qn("w:tc")) for tr in table.findall(qn("w:tr"))]
    assert el_text(cells[0][0]).replace(" ", "") == "学校" and el_text(cells[1][0]) == "参赛队号"
    def write(tc, text, bottom=False):
        if bottom:
            tcpr = tc.get_or_add_tcPr()
            va = tcpr.find(qn("w:vAlign"))
            if va is None:
                va = OxmlElement("w:vAlign"); tcpr.append(va)
            va.set(qn("w:val"), "bottom")
        p = tc.findall(qn("w:p"))[-1]
        runs = p.findall(qn("w:r")) or [p.makeelement(qn("w:r"), {})]
        if runs[-1].getparent() is None:
            p.append(runs[-1])
        t = runs[-1].find(qn("w:t"))
        if t is None:
            t = OxmlElement("w:t"); runs[-1].append(t)
        t.text = (t.text or "") + text
        t.set(qn("xml:space"), "preserve")
    write(cells[0][1], SCHOOL, bottom=True)
    if team_no:
        write(cells[1][1], team_no, bottom=True)
    for k, name in enumerate(MEMBERS):
        assert el_text(cells[2 + k][1]) == f"{k + 1}."
        write(cells[2 + k][1], " " + name)


def add_cover(doc, team_no: str) -> None:
    """模板第0页（徽标、赛事名称、信息表）插在摘要页之前；摘要页另起一页，页码自摘要页起为1。"""
    tpl = Document(TEMPLATE)
    tk = list(tpl.element.body.iterchildren())
    assert el_text(tk[2]) == "中国研究生创新实践系列大赛" and tk[6].tag == qn("w:tbl")
    body = doc.element.body
    first = body[0]
    assert el_text(first) == "中国研究生创新实践系列大赛", "提交版首段应为摘要页抬头"
    cover = []
    for i in range(7):
        el = copy.deepcopy(tk[i]); normalize_ooxml(el); cover.append(el)
    A = "http://schemas.openxmlformats.org/drawingml/2006/main"
    R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    for blip in cover[0].iter(f"{{{A}}}blip"):
        rid, _ = doc.part.get_or_add_image(io.BytesIO(tpl.part.related_parts[blip.get(f"{{{R}}}embed")].blob))
        blip.set(f"{{{R}}}embed", rid)
    for j, dp in enumerate(cover[0].iter(qn("wp:docPr"))):
        dp.set("id", str(20001 + j))
    fill_cover(cover[6], team_no)
    for el in cover:
        first.addprevious(el)
    Paragraph(first, doc._body).paragraph_format.page_break_before = True
    sec = doc.sections[-1]
    sec.different_first_page_header_footer = True            # 封面用空白首页页脚，不显示页码
    for part in (sec.first_page_footer,):
        for p in part.paragraphs:
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
    sp = sec._sectPr
    for el in list(sp):                                        # 页眉全为空：删去引用，彻底“无页眉”；未启用奇偶页，偶数页页脚引用无用
        if el.tag == qn("w:headerReference") or (el.tag == qn("w:footerReference") and el.get(qn("w:type")) == "even"):
            sp.remove(el)
    pg = sp.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType"); sp.append(pg)
    pg.set(qn("w:start"), "0")


LOCAL_PREFIXES = ["/home/dreamboat"]            # 实验机本机用户目录；提交包内统一替换，避免评阅时暴露身份


def sanitize_support(root: Path, zpath: Path) -> int:
    """把归档记录中的本机绝对路径前缀换成<实验机目录>，重写manifest中的文件哈希并重新打包。仓库原件不改。"""
    import zipfile
    n = 0
    for p in sorted(root.rglob("*")):
        if not p.is_file() or p.suffix not in (".json", ".csv", ".md", ".py", ".txt", ".log"):
            continue
        t = p.read_text(encoding="utf-8", errors="strict")
        new = t
        for pre in LOCAL_PREFIXES:
            new = new.replace(pre, "<实验机目录>")
        if new != t:
            p.write_text(new, encoding="utf-8"); n += 1
    man = json.loads((root / "manifest.json").read_text(encoding="utf-8"))
    man["path_sanitization"] = {"replaced_prefix": "本机用户目录", "with": "<实验机目录>", "files": n,
                                "note": "仅替换路径文字，数值与字段不变；仓库原件保留原路径"}
    man["files"] = [{"path": q.relative_to(root).as_posix(), "bytes": q.stat().st_size, "sha256": sha(q)}
                    for q in sorted(root.rglob("*")) if q.is_file() and q.name != "manifest.json"]
    (root / "manifest.json").write_text(json.dumps(man, ensure_ascii=False, indent=2), encoding="utf-8")
    for q in root.rglob("*"):
        if q.is_file():
            assert all(pre not in q.read_text(encoding="utf-8", errors="ignore") for pre in LOCAL_PREFIXES), q
    zpath.unlink()
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
        for q in sorted(root.rglob("*")):
            if q.is_file():
                z.write(q, Path("支撑材料") / q.relative_to(root))
    return n


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--team-no", default="", help="参赛队号；为空时封面该栏留空")
    a = ap.parse_args()
    assert sha(SRC_DOCX) == SRC_SHA, "LBSS提交版底稿已变化，停止并核对"
    doc = Document(SRC_DOCX)
    counts = apply_edits(doc)                                  # 符号表中的“只计”属于预定润色，故在润色后取表格快照
    before = [el_text(t) for t in doc.element.body.iter(qn("w:tbl"))]
    add_cover(doc, a.team_no)
    after = [el_text(t) for t in doc.element.body.iter(qn("w:tbl"))][1:]   # 去掉新加的封面信息表
    assert before == after, "表格文字发生变化"
    reorder_children(doc)
    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)

    # 支撑材料：沿用LBSS轮的打包函数，结构与论文附录B、C的路径一致
    sys.path.insert(0, str(SUB))
    from support import package_support
    stage = HERE / "_stage"
    shutil.rmtree(stage, ignore_errors=True)
    stats = json.loads((SUB / "out/补报统计.json").read_text(encoding="utf-8"))
    info = package_support(R12, stage, stats)
    sanitized = sanitize_support(stage / "out/支撑材料", stage / "out/支撑材料.zip")
    dst = HERE / "02_支撑材料"
    shutil.rmtree(dst, ignore_errors=True)
    dst.mkdir(parents=True)
    shutil.move(str(stage / "out/支撑材料.zip"), dst / "支撑材料.zip")
    (HERE / "03_工具").mkdir(exist_ok=True)
    shutil.move(str(stage / "out/Windows_Word正式导出.ps1"), HERE / "03_工具/Windows_Word正式导出.ps1")
    shutil.rmtree(stage)
    receipt = {"source_docx_sha256": SRC_SHA, "edits": counts, "cover": {"school": SCHOOL, "team_no": a.team_no, "members": MEMBERS},
               "tables_text_unchanged": True, "submission_docx_sha256": sha(OUT_DOCX), "support": info,
               "support_zip_sha256": sha(dst / "支撑材料.zip"), "path_sanitized_files": sanitized, "solver_or_evaluator_executed": False}
    (HERE / "04_核验记录").mkdir(exist_ok=True)
    (HERE / "04_核验记录/prepare_receipt.json").write_text(json.dumps(receipt, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: v for k, v in receipt.items() if k != "support"}, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
