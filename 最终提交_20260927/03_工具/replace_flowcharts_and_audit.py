#!/usr/bin/env python3
# AI-assisted maintenance utility. It does not modify experiments or numerical results.\n# Triggered only for the final-submission document review branch.\n# Figure lookup scans all document-order paragraphs, including table-contained objects.
from __future__ import annotations
import argparse, hashlib, json, os, re, shutil, tempfile, zipfile
from pathlib import Path
from xml.etree import ElementTree as ET
from PIL import Image
from docx import Document

NS = {
    "w":"http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "a":"http://schemas.openxmlformats.org/drawingml/2006/main",
    "r":"http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "pr":"http://schemas.openxmlformats.org/package/2006/relationships",
    "v":"urn:schemas-microsoft-com:vml",
}
CAPTIONS = {
    "图2-3":"F35.png",
    "图5-4":"F38.png",
    "图6-4":"F39.png",
    "图7-2":"F40.png",
}
STALE_TERMS = [
    "原版A评价","原版B完整评价","原版 Step 2","原版Step 2",
    "近优门控","两字段方案","不补0","不补 0","不删图","字典序定案",
    "C3","R12","R22","round12","F27","F28","A087",
]
def sha256(p: Path)->str:
    h=hashlib.sha256()
    with p.open("rb") as f:
        for b in iter(lambda:f.read(1<<20),b""): h.update(b)
    return h.hexdigest()

def para_text(p):
    return "".join((t.text or "") for t in p.findall(".//w:t",NS))

def paragraph_image_rid(p):
    blip=p.find(".//a:blip",NS)
    if blip is not None:
        rid=blip.get("{%s}embed"%NS["r"])
        if rid:return rid
    vm=p.find(".//v:imagedata",NS)
    if vm is not None:
        rid=vm.get("{%s}id"%NS["r"])
        if rid:return rid
    return None

def nearest_image_rid(paras, idx):
    # Captions in this submission are normally below figures. Prefer backward search.
    for j in range(idx-1, max(-1,idx-60), -1):
        rid=paragraph_image_rid(paras[j])
        if rid:return j,rid
    # Compatibility/layout edits may place the caption before an anchored image.
    for j in range(idx+1, min(len(paras),idx+25)):
        rid=paragraph_image_rid(paras[j])
        if rid:return j,rid
    ctx=[{"i":j,"text":para_text(paras[j])[:120],"rid":paragraph_image_rid(paras[j])}
         for j in range(max(0,idx-12),min(len(paras),idx+13))]
    raise RuntimeError(f"No nearby image for caption paragraph {idx}; context={ctx}")

def replace_images(src:Path,dst:Path,flowdir:Path):
    with tempfile.TemporaryDirectory() as td:
        td=Path(td)
        with zipfile.ZipFile(src) as z:z.extractall(td)
        docxml=td/"word/document.xml"
        relxml=td/"word/_rels/document.xml.rels"
        tree=ET.parse(docxml); root=tree.getroot()
        body=root.find("w:body",NS)
        paras=body.findall(".//w:p",NS)
        reltree=ET.parse(relxml); relroot=reltree.getroot()
        relmap={x.get("Id"):x.get("Target") for x in relroot}
        records=[]
        used=set()
        for cap,source_name in CAPTIONS.items():
            matches=[i for i,p in enumerate(paras) if re.match(r"^\\s*"+re.escape(cap)+r"(?:\\s|$)", para_text(p))]
            if not matches: raise RuntimeError(f"Caption not found: {cap}")
            # Match the actual caption paragraph only; prose such as "见图5-4" must never qualify.
            idx=matches[0]
            pidx,rid=nearest_image_rid(paras,idx)
            if rid in used: raise RuntimeError(f"Relationship reused unexpectedly: {rid}")
            used.add(rid)
            target=relmap.get(rid)
            if not target: raise RuntimeError(f"Relationship target missing: {rid}")
            media=(td/"word"/target).resolve()
            if not str(media).startswith(str((td/"word").resolve())): raise RuntimeError("unsafe target")
            srcimg=flowdir/source_name
            if not srcimg.exists(): raise RuntimeError(f"New image missing: {srcimg}")
            ext=media.suffix.lower()
            with Image.open(srcimg) as im:
                new_wh=im.size
                if ext in (".jpg",".jpeg"):
                    im.convert("RGB").save(media,"JPEG",quality=98,subsampling=0)
                elif ext==".png":
                    shutil.copy2(srcimg,media)
                else:
                    raise RuntimeError(f"Unsupported embedded image type {ext} for {cap}: {media}")
            try:
                with Image.open(media) as im2: final_wh=im2.size
            except Exception: final_wh=None
            records.append({
                "caption":cap,"caption_para_index":idx,"image_para_index":pidx,
                "relationship_id":rid,"target":target,"new_source":str(srcimg),
                "new_source_sha256":sha256(srcimg),"embedded_sha256":sha256(media),
                "new_source_px":list(new_wh),"embedded_px":list(final_wh) if final_wh else None,
            })
        # Rezip deterministically enough for review; Word content/order retained.
        with zipfile.ZipFile(dst,"w",zipfile.ZIP_DEFLATED) as z:
            for p in sorted(td.rglob("*")):
                if p.is_file(): z.write(p,p.relative_to(td).as_posix())
    return records

def extract_text(docx:Path):
    d=Document(docx)
    lines=[]
    for p in d.paragraphs:
        t=p.text.strip()
        if t: lines.append(t)
    for ti,tbl in enumerate(d.tables):
        for ri,row in enumerate(tbl.rows):
            rowtxt=" | ".join(c.text.strip().replace("\n"," / ") for c in row.cells)
            if rowtxt.strip(" |"): lines.append(f"[TABLE {ti+1} ROW {ri+1}] {rowtxt}")
    return "\n".join(lines)

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--src",required=True);ap.add_argument("--dst",required=True)
    ap.add_argument("--flowdir",required=True);ap.add_argument("--report",required=True)
    ap.add_argument("--text",required=True)
    args=ap.parse_args()
    src,dst,flowdir=map(Path,[args.src,args.dst,args.flowdir])
    rec=replace_images(src,dst,flowdir)
    txt=extract_text(dst);Path(args.text).write_text(txt,encoding="utf-8")
    stale={t:len(re.findall(re.escape(t),txt,re.I)) for t in STALE_TERMS}
    stale={k:v for k,v in stale.items() if v}
    title_lines=[x for x in txt.splitlines() if ("题 目" in x or "多核NPU" in x or "多核 NPU" in x)][:30]
    report={
      "source":str(src),"source_sha256":sha256(src),"output":str(dst),"output_sha256":sha256(dst),
      "replacements":rec,"replacement_count":len(rec),"stale_terms":stale,
      "title_candidates":title_lines,
      "scope":"Only four embedded flowchart images replaced; no prose, equations, tables, results, headers, footers, or section settings intentionally edited."
    }
    Path(args.report).write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(report,ensure_ascii=False,indent=2))
    assert len(rec)==4
if __name__=="__main__":main()
